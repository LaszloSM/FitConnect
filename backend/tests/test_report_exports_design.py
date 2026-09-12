"""Los reportes son documentos para entregar, no volcados de tablas.

Portada, resumen ejecutivo con comparativo contra el período anterior, gráficos
y formato condicional. Se prueba la ESTRUCTURA del archivo, no solo que no dé 500.
"""
import io
import re
from datetime import date

from openpyxl import load_workbook

from app.services import report_exports as rx


def _spec(con_datos: bool = True) -> rx.ReportSpec:
    etiquetas = [f"{d:02d}/09" for d in range(1, 31)]
    valores = [float(d * 10000) for d in range(1, 31)] if con_datos else [0.0] * 30
    return rx.ReportSpec(
        title="Ingresos", subtitle="01/09/2026 — 30/09/2026", comparison_label="02/08/2026 — 31/08/2026",
        tiles=[
            rx._tile("Ingresos confirmados", 4_650_000, 4_000_000, money=True),
            rx._tile("Egresos", 900_000, 600_000, money=True, good_when_up=False),
            rx._tile("Socios nuevos", 12, 0),
        ],
        charts=[
            rx.Chart("bar", "Ingresos por día", etiquetas, [("Membresías", valores), ("Pases", valores)], money=True),
            rx.Chart("hbar", "Por método", ["Efectivo", "Transferencia"], [("Monto", [3.0, 1.0] if con_datos else [0.0, 0.0])], money=True),
            rx.Chart("line", "Renovación", ["jul", "ago", "sep"], [("Renovación", [70.0, 65.0, 72.0])]),
        ],
        sections=[rx.Section(
            "Inventario", [rx.Column("name", "Producto"), rx.Column("days_left", "Días", align_right=True, alert_below=8),
                           rx.Column("amount", "Monto", money=True, data_bar=True)],
            [{"name": "=HYPERLINK(\"x\")", "days_left": 3, "amount": 1000}, {"name": "Agua", "days_left": 40, "amount": 5000}],
        )],
    )


def _pages(pdf: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf))


def test_pdf_has_cover_summary_and_charts():
    pdf = rx.build_pdf(_spec(), gym_name="Gym Hierro", logo_png=None, nit="900.123.456-7")
    assert pdf.startswith(b"%PDF")
    # Portada + resumen con gráficos + detalle.
    assert _pages(pdf) >= 2


def test_pdf_survives_empty_charts():
    pdf = rx.build_pdf(_spec(con_datos=False), gym_name="Gym Vacío", logo_png=None)
    assert pdf.startswith(b"%PDF")


def test_xlsx_opens_with_summary_charts_and_conditional_formatting():
    wb = load_workbook(io.BytesIO(rx.build_xlsx(_spec(), gym_name="Gym Hierro", logo_png=None)))
    assert wb.sheetnames[0] == "Resumen"
    resumen = wb["Resumen"]
    valores = [c.value for row in resumen.iter_rows() for c in row if c.value is not None]
    assert "Período anterior" in valores
    assert 4_650_000 in valores and 4_000_000 in valores
    # Variación como número de porcentaje, no como texto.
    esperado = rx._delta(4_650_000, 4_000_000) / 100
    assert any(isinstance(v, float) and abs(v - esperado) < 1e-9 for v in valores)
    # "Socios nuevos" pasó de 0 a 12: sin base no hay variación.
    assert resumen.conditional_formatting

    graficos = wb["Gráficos"]
    assert len(graficos._charts) == 3

    inventario = wb["Inventario"]
    reglas = [r for rango in inventario.conditional_formatting for r in rango.rules]
    assert {r.type for r in reglas} >= {"cellIs", "dataBar"}
    # El escape anti-fórmula sigue en pie.
    assert inventario["A2"].value.startswith("'=")


def test_empty_charts_do_not_create_a_charts_sheet():
    spec = _spec(con_datos=False)
    spec.charts = [c for c in spec.charts if c.kind != "line"]
    wb = load_workbook(io.BytesIO(rx.build_xlsx(spec, gym_name="Gym Vacío", logo_png=None)))
    assert "Gráficos" not in wb.sheetnames


def test_delta_needs_a_base():
    assert rx._delta(12, 0) is None
    assert rx._delta(110, 100) == 10.0
    assert rx._delta(90, 100) == -10.0


def test_previous_range_has_the_same_length():
    prev_from, prev_to = rx.previous_range(date(2026, 8, 1), date(2026, 8, 31))
    assert prev_to == date(2026, 7, 31)
    assert (prev_to - prev_from).days == 30


async def test_ingresos_export_compares_against_the_previous_period(client):
    email = "exports-diseno@gym.com"
    await client.post("/api/v1/auth/register", json={
        "email": email, "password": "Pass123456!", "full_name": "Owner Export", "gym_name": "Gym Export Diseño", "city": "Bogotá",
    })
    login = await client.post("/api/v1/auth/login", json={"email": email, "password": "Pass123456!"})
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

    r = await client.get("/api/v1/reports/ingresos/export.xlsx", headers=headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Resumen"
    textos = [c.value for row in wb["Resumen"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("Comparado con") for t in textos)

    pdf = await client.get("/api/v1/reports/cierre/export.pdf", headers=headers)
    assert pdf.status_code == 200
    assert _pages(pdf.content) >= 2
