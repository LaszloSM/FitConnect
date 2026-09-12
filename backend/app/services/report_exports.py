"""Exports por pestaña (PDF y Excel) con la marca del gym.

Un `ReportSpec` describe el documento: título, período, indicadores con su
comparativo contra el período anterior, gráficos y secciones (tabla + totales).
Los mismos datasets alimentan los dos formatos.

Un reporte que solo trae tablas no sirve para decidir: el dueño necesita saber
primero si el mes fue mejor o peor que el anterior, y después el detalle. Por
eso el PDF abre con portada y resumen ejecutivo, y el Excel con una hoja
"Resumen" y gráficos nativos que siguen vivos al filtrar.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, timedelta
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    CondPageBreak, Image as PdfImage, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
    TableStyle,
)
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.timeutil import today_bogota
from app.services import report_data as rd
from app.services.exports import _BRAND, _BRAND_TEXT, _BRAND_TINT, _INK, _MONEY_FMT, _cop


@dataclass
class Column:
    key: str
    header: str
    money: bool = False
    align_right: bool = False
    width_mm: float | None = None
    # Formato condicional en Excel: lo que merece mirarse primero.
    alert_below: float | None = None
    alert_above: float | None = None
    data_bar: bool = False


@dataclass
class Section:
    title: str
    columns: list[Column]
    rows: list[dict]
    total_row: dict | None = None
    note: str | None = None
    empty_text: str = "Sin datos en el período."


@dataclass
class Tile:
    label: str
    value: str
    previous: str | None = None
    delta_pct: float | None = None
    # En egresos o denegados, subir es malo.
    good_when_up: bool = True
    raw: float | None = None
    raw_previous: float | None = None
    money: bool = False


@dataclass
class Chart:
    kind: str  # "bar" | "hbar" | "line"
    title: str
    labels: list[str]
    series: list[tuple[str, list[float]]]
    money: bool = False


@dataclass
class ReportSpec:
    title: str
    subtitle: str
    tiles: list[Tile] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)
    charts: list[Chart] = field(default_factory=list)
    comparison_label: str | None = None


MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

_OK = "15803D"
_BAD = "B91C1C"
_SERIES = [_BRAND, "57534E", "A8A29E"]
_MUTED = "78716C"
_SURFACE = "F7F5F3"


def _xlsx_safe(v):
    """Excel interpreta '=', '+', '-', '@' iniciales como fórmula: se neutraliza
    el texto que viene de nombres/conceptos escritos por el usuario."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _fmt(value, money: bool) -> str:
    if value is None or value == "":
        return "—"
    if money:
        return _cop(float(value))
    if isinstance(value, float):
        return f"{value:,.1f}".replace(",", ".")
    return str(value)


def _short_cop(v: float) -> str:
    a = abs(v)
    if a >= 1_000_000:
        txt = f"{v / 1_000_000:.1f}".replace(".", ",")
        return f"${txt[:-2] if txt.endswith(',0') else txt} M"
    if a >= 1_000:
        return f"${v / 1_000:.0f} mil"
    return f"${v:.0f}"


def _delta(current: float | None, previous: float | None) -> float | None:
    """Sin base no hay variación: pasar de 0 a algo no es "+∞%"."""
    if current is None or previous is None or previous == 0:
        return None
    return round((current - previous) / abs(previous) * 100, 1)


def _tile(label: str, current: float, previous: float | None, *, money: bool = False, good_when_up: bool = True) -> Tile:
    show = _cop if money else (lambda v: f"{v:,.0f}".replace(",", "."))
    return Tile(
        label=label, value=show(current),
        previous=show(previous) if previous is not None else None,
        delta_pct=_delta(current, previous), good_when_up=good_when_up,
        raw=float(current), raw_previous=float(previous) if previous is not None else None, money=money,
    )


def previous_range(from_date: date | None, to_date: date | None) -> tuple[date, date]:
    """El período inmediatamente anterior con la misma cantidad de días."""
    start_day, end_day, _, _ = rd.period_bounds(from_date, to_date)
    prev_end = start_day - timedelta(days=1)
    return prev_end - (end_day - start_day), prev_end


def _period_label(from_day: str | date, to_day: str | date) -> str:
    f = (date.fromisoformat(from_day) if isinstance(from_day, str) else from_day).strftime("%d/%m/%Y")
    t = (date.fromisoformat(to_day) if isinstance(to_day, str) else to_day).strftime("%d/%m/%Y")
    return f"{f} — {t}" if f != t else f


def _highlights(tiles: list[Tile]) -> list[str]:
    """Las tres variaciones más grandes, dichas en una frase."""
    con_delta = sorted((t for t in tiles if t.delta_pct is not None), key=lambda t: abs(t.delta_pct), reverse=True)
    frases = []
    for t in con_delta[:3]:
        if t.delta_pct == 0:
            continue
        sentido = "subió" if t.delta_pct > 0 else "bajó"
        frases.append(f"{t.label}: {sentido} {abs(t.delta_pct):.1f}% frente al período anterior (de {t.previous} a {t.value}).".replace(".0%", "%"))
    return frases


# ─── PDF ───

def _visible_series(chart: Chart) -> list[tuple[str, list[float]]]:
    return [(name, serie) for name, serie in chart.series if any(v > 0 for v in serie)]


def _chart_drawing(chart: Chart, width: float, height: float) -> Drawing | None:
    chart = Chart(chart.kind, chart.title, chart.labels, _visible_series(chart), chart.money)
    valores = [v for _, serie in chart.series for v in serie]
    if not chart.labels or not valores or max(valores) <= 0:
        return None
    if chart.kind == "hbar":
        # Una barra sola no puede ocupar media página.
        height = min(height, 16 * mm + 9 * mm * len(chart.labels))
    d = Drawing(width, height)
    fmt = _short_cop if chart.money else (lambda v: f"{v:,.0f}".replace(",", "."))
    n = len(chart.labels)
    # Con muchas categorías se rotulan algunas: 31 fechas encimadas no se leen.
    paso = max(1, -(-n // 12))
    etiquetas = [lbl if i % paso == 0 else "" for i, lbl in enumerate(chart.labels)]

    if chart.kind == "hbar":
        c = HorizontalBarChart()
        c.x, c.y, c.width, c.height = 48 * mm, 6 * mm, width - 56 * mm, height - 12 * mm
        c.data = [tuple(reversed(s)) for _, s in chart.series]
        c.categoryAxis.categoryNames = [lbl[:26] for lbl in reversed(chart.labels)]
        c.categoryAxis.labels.boxAnchor = "e"
        c.categoryAxis.labels.dx = -3
    elif chart.kind == "line":
        c = HorizontalLineChart()
        c.x, c.y, c.width, c.height = 16 * mm, 10 * mm, width - 22 * mm, height - 18 * mm
        c.data = [tuple(s) for _, s in chart.series]
        c.categoryAxis.categoryNames = etiquetas
        c.joinedLines = 1
        for i in range(len(chart.series)):
            c.lines[i].strokeColor = colors.HexColor("#" + _SERIES[i % len(_SERIES)])
            c.lines[i].strokeWidth = 1.8
    else:
        c = VerticalBarChart()
        c.x, c.y, c.width, c.height = 16 * mm, 10 * mm, width - 22 * mm, height - 18 * mm
        c.data = [tuple(s) for _, s in chart.series]
        c.categoryAxis.categoryNames = etiquetas
        c.groupSpacing = 3 if n > 20 else 6
        c.barSpacing = 1

    if chart.kind != "line":
        for i in range(len(chart.series)):
            c.bars[i].fillColor = colors.HexColor("#" + _SERIES[i % len(_SERIES)])
            c.bars[i].strokeColor = None
    c.valueAxis.valueMin = 0
    c.valueAxis.labelTextFormat = fmt
    c.valueAxis.labels.fontName = "Helvetica"
    c.valueAxis.labels.fontSize = 6.5
    c.valueAxis.labels.fillColor = colors.HexColor("#" + _MUTED)
    c.valueAxis.visibleGrid = 1
    c.valueAxis.gridStrokeColor = colors.HexColor("#E7E5E4")
    c.valueAxis.strokeColor = colors.HexColor("#D6D3D1")
    c.categoryAxis.labels.fontName = "Helvetica"
    c.categoryAxis.labels.fontSize = 6.5
    c.categoryAxis.labels.fillColor = colors.HexColor("#" + _MUTED)
    c.categoryAxis.strokeColor = colors.HexColor("#D6D3D1")
    d.add(c)

    if len(chart.series) > 1:
        legend = Legend()
        legend.x, legend.y = width - 4 * mm, height - 2 * mm
        legend.alignment = "right"
        legend.boxAnchor = "ne"
        legend.fontName = "Helvetica"
        legend.fontSize = 7
        legend.columnMaximum = 1
        legend.deltax = 70
        legend.dxTextSpace = 4
        legend.fillColor = colors.HexColor("#" + _MUTED)
        legend.colorNamePairs = [
            (colors.HexColor("#" + _SERIES[i % len(_SERIES)]), name) for i, (name, _) in enumerate(chart.series)
        ]
        d.add(legend)
    return d


def build_pdf(spec: ReportSpec, *, gym_name: str, logo_png: bytes | None, nit: str | None = None) -> bytes:
    buf = io.BytesIO()
    margin = 16 * mm
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=22 * mm, leftMargin=margin, rightMargin=margin,
                            title=f"{spec.title} · {gym_name}", author=gym_name, creator="FitConnect")
    avail = A4[0] - 2 * margin
    styles = getSampleStyleSheet()
    ink = colors.HexColor("#" + _INK)
    muted = colors.HexColor("#" + _MUTED)
    brand = colors.HexColor("#" + _BRAND)
    base = ParagraphStyle("base", parent=styles["Normal"], fontName="Helvetica", textColor=ink)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=21, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=12, leading=15, spaceBefore=8, spaceAfter=4)
    note_style = ParagraphStyle("n", parent=base, fontSize=8, leading=10, textColor=muted)
    cell = ParagraphStyle("c", parent=base, fontSize=8.5, leading=10.5)
    generated = today_bogota().strftime("%d/%m/%Y")

    def _cover(canvas, _doc):
        canvas.saveState()
        canvas.setFillColor(brand)
        canvas.rect(0, 0, 9 * mm, A4[1], stroke=0, fill=1)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawString(margin + 6 * mm, 14 * mm, f"Generado el {generated} con FitConnect")
        canvas.restoreState()

    def _page(canvas, _doc):
        canvas.saveState()
        canvas.setFillColor(brand)
        canvas.rect(0, A4[1] - 4 * mm, A4[0], 4 * mm, stroke=0, fill=1)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(margin, 12 * mm, f"{gym_name} · {spec.title} · {spec.subtitle}")
        canvas.drawRightString(A4[0] - margin, 12 * mm, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    els: list = [Spacer(1, 38 * mm)]

    # ── Portada ──
    if logo_png:
        try:
            iw, ih = ImageReader(io.BytesIO(logo_png)).getSize()
            w = min(40 * mm, 26 * mm * iw / ih)
            els += [PdfImage(io.BytesIO(logo_png), width=w, height=w * ih / iw, hAlign="LEFT"), Spacer(1, 6 * mm)]
        except Exception:
            pass
    els.append(Paragraph(escape(gym_name), ParagraphStyle("gym", parent=base, fontName="Helvetica-Bold", fontSize=22, leading=26)))
    if nit:
        els.append(Paragraph(f"NIT {escape(nit)}", ParagraphStyle("nit", parent=base, fontSize=10, textColor=muted)))
    els += [
        Spacer(1, 42 * mm),
        Paragraph("REPORTE", ParagraphStyle("eyebrow", parent=base, fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#" + _BRAND_TEXT))),
        Spacer(1, 2 * mm),
        Paragraph(escape(spec.title), ParagraphStyle("ct", parent=base, fontName="Helvetica-Bold", fontSize=34, leading=38)),
        Spacer(1, 3 * mm),
        Paragraph(escape(spec.subtitle), ParagraphStyle("cs", parent=base, fontSize=13, leading=16, textColor=muted)),
    ]
    if spec.comparison_label:
        els.append(Paragraph(f"Comparado con {escape(spec.comparison_label)}", ParagraphStyle("cc", parent=base, fontSize=10, leading=14, textColor=muted)))
    els.append(PageBreak())

    # ── Resumen ejecutivo ──
    els.append(Paragraph("Resumen ejecutivo", h1))
    frases = _highlights(spec.tiles)
    for f in frases:
        els.append(Paragraph(escape(f), ParagraphStyle("hl", parent=base, fontSize=9.5, leading=13, bulletIndent=0, leftIndent=8), bulletText="·"))
    if frases:
        els.append(Spacer(1, 4 * mm))

    if spec.tiles:
        per_row = min(len(spec.tiles), 4)
        label_s = ParagraphStyle("tl", parent=base, fontSize=7.5, leading=9, textColor=muted)
        value_s = ParagraphStyle("tv", parent=base, fontName="Helvetica-Bold", fontSize=14, leading=17)
        for i in range(0, len(spec.tiles), per_row):
            chunk = spec.tiles[i:i + per_row]
            fila = []
            for t in chunk:
                partes = [Paragraph(escape(t.label), label_s), Paragraph(escape(t.value), value_s)]
                if t.delta_pct is not None:
                    bueno = (t.delta_pct >= 0) == t.good_when_up
                    color = _OK if bueno or t.delta_pct == 0 else _BAD
                    signo = "+" if t.delta_pct > 0 else ""
                    partes.append(Paragraph(
                        f'<font color="#{color}"><b>{signo}{t.delta_pct:.1f}%</b></font> vs. {escape(t.previous or "—")}'.replace(".0%", "%"),
                        ParagraphStyle("td", parent=base, fontSize=7.5, leading=9, textColor=muted),
                    ))
                elif t.previous is not None:
                    partes.append(Paragraph(f"Antes: {escape(t.previous)}", ParagraphStyle("tp", parent=base, fontSize=7.5, leading=9, textColor=muted)))
                fila.append(partes)
            tabla = Table([fila], colWidths=[avail / per_row] * len(chunk), hAlign="LEFT")
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + _SURFACE)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBEFORE", (1, 0), (-1, -1), 3, colors.white),
                ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ]))
            els += [tabla, Spacer(1, 2 * mm)]
        els.append(Spacer(1, 4 * mm))

    for ch in spec.charts:
        drawing = _chart_drawing(ch, avail, 62 * mm)
        if drawing is None:
            continue
        els.append(KeepTogether([Paragraph(escape(ch.title), h2), drawing, Spacer(1, 3 * mm)]))

    # ── Detalle ──
    for sec in spec.sections:
        titulo = Paragraph(escape(sec.title), h2)
        if not sec.rows:
            els.append(KeepTogether([titulo, Paragraph(escape(sec.empty_text), note_style), Spacer(1, 3 * mm)]))
            continue
        data = [[c.header for c in sec.columns]]
        for r in sec.rows:
            data.append([
                _fmt(r.get(c.key), c.money) if (c.money or c.align_right) else Paragraph(escape(_fmt(r.get(c.key), c.money)), cell)
                for c in sec.columns
            ])
        if sec.total_row:
            data.append([_fmt(sec.total_row.get(c.key), c.money) if c.key in sec.total_row else "" for c in sec.columns])
        fixed = sum(c.width_mm for c in sec.columns if c.width_mm)
        flex = [c for c in sec.columns if not c.width_mm]
        widths = [c.width_mm * mm if c.width_mm else max((avail - fixed * mm) / max(len(flex), 1), 20 * mm) for c in sec.columns]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + _BRAND_TINT)),
            ("TEXTCOLOR", (0, 0), (-1, 0), ink),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAF9F7")]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E7E5E4")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for i, c in enumerate(sec.columns):
            if c.money or c.align_right:
                style.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
        if sec.total_row:
            style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("LINEABOVE", (0, -1), (-1, -1), 1, ink),
                      ("BACKGROUND", (0, -1), (-1, -1), colors.white)]
        t.setStyle(TableStyle(style))
        bloque = [titulo, t] + ([Paragraph(escape(sec.note), note_style)] if sec.note else [])
        if len(sec.rows) <= 12:
            els.append(KeepTogether(bloque))
        else:
            # Que el título no quede huérfano al pie de la página.
            els += [CondPageBreak(40 * mm), *bloque]
        els.append(Spacer(1, 3 * mm))

    doc.build(els, onFirstPage=_cover, onLaterPages=_page)
    return buf.getvalue()


# ─── Excel ───

_SHEET_BAD = str.maketrans({c: " " for c in "[]:*?/\\"})


def _sheet_title(title: str, used: set[str]) -> str:
    base = title.translate(_SHEET_BAD).strip()[:28] or "Hoja"
    nombre, k = base, 2
    while nombre in used:
        nombre = f"{base[:25]} {k}"
        k += 1
    used.add(nombre)
    return nombre


def _header(ws, row: int, values: list[str], start_col: int = 1) -> None:
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font = Font(bold=True, color=_INK)
        c.fill = PatternFill("solid", fgColor=_BRAND_TINT)
        c.alignment = Alignment(horizontal="center")


def build_xlsx(spec: ReportSpec, *, gym_name: str, logo_png: bytes | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    used = {"Resumen"}

    row = 2
    if logo_png:
        try:
            img = XlsxImage(io.BytesIO(logo_png))
            scale = min(1.0, 120 / max(img.width, 1), 60 / max(img.height, 1))
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            ws.add_image(img, "B2")
            row = 6
        except Exception:
            pass
    ws.cell(row=row, column=2, value=gym_name).font = Font(bold=True, size=18, color=_INK)
    ws.cell(row=row + 1, column=2, value=spec.title).font = Font(bold=True, size=13, color=_BRAND_TEXT)
    ws.cell(row=row + 2, column=2, value=spec.subtitle).font = Font(size=10, color=_MUTED)
    if spec.comparison_label:
        ws.cell(row=row + 3, column=2, value=f"Comparado con {spec.comparison_label}").font = Font(size=10, color=_MUTED)
    ws.cell(row=row + 4, column=2, value=f"Generado el {today_bogota().strftime('%d/%m/%Y')} · FitConnect").font = Font(size=9, color="A8A29E")

    r = row + 6
    if spec.tiles:
        _header(ws, r, ["Indicador", "Período", "Período anterior", "Variación"], start_col=2)
        for t in spec.tiles:
            r += 1
            ws.cell(row=r, column=2, value=t.label)
            actual = ws.cell(row=r, column=3, value=t.raw if t.raw is not None else _xlsx_safe(t.value))
            antes = ws.cell(row=r, column=4, value=t.raw_previous if t.raw_previous is not None else (t.previous or "—"))
            if t.money:
                actual.number_format = antes.number_format = _MONEY_FMT
            elif t.raw is not None:
                actual.number_format = antes.number_format = "#,##0"
            var = ws.cell(row=r, column=5, value=(t.delta_pct / 100) if t.delta_pct is not None else "—")
            if t.delta_pct is not None:
                var.number_format = "+0.0%;-0.0%;0.0%"
                ref = f"E{r}"
                sube, baja = (_OK, _BAD) if t.good_when_up else (_BAD, _OK)
                ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], font=Font(bold=True, color=sube)))
                ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, color=baja)))
        frases = _highlights(spec.tiles)
        if frases:
            r += 2
            ws.cell(row=r, column=2, value="Lo más importante").font = Font(bold=True, color=_INK)
            for f in frases:
                r += 1
                ws.cell(row=r, column=2, value=f).font = Font(color="44403C")
    for col, width in (("A", 3), ("B", 42), ("C", 18), ("D", 18), ("E", 14)):
        ws.column_dimensions[col].width = width

    graficables = [
        Chart(c.kind, c.title, c.labels, _visible_series(c), c.money)
        for c in spec.charts if c.labels and _visible_series(c)
    ]
    if graficables:
        wg = wb.create_sheet(_sheet_title("Gráficos", used))
        wg.sheet_view.showGridLines = False
        fila = 1
        for ch in graficables:
            wg.cell(row=fila, column=1, value=ch.title).font = Font(bold=True, size=12, color=_INK)
            head = fila + 1
            _header(wg, head, ["Categoría", *[name for name, _ in ch.series]])
            for i, label in enumerate(ch.labels):
                wg.cell(row=head + 1 + i, column=1, value=_xlsx_safe(label))
                for j, (_, serie) in enumerate(ch.series):
                    c = wg.cell(row=head + 1 + i, column=2 + j, value=serie[i])
                    if ch.money:
                        c.number_format = _MONEY_FMT
            ultimo = head + len(ch.labels)
            grafico = LineChart() if ch.kind == "line" else BarChart()
            if ch.kind == "hbar":
                grafico.type = "bar"
            grafico.title = ch.title
            grafico.height, grafico.width = 7.5, 17
            grafico.add_data(Reference(wg, min_col=2, max_col=1 + len(ch.series), min_row=head, max_row=ultimo), titles_from_data=True)
            grafico.set_categories(Reference(wg, min_col=1, min_row=head + 1, max_row=ultimo))
            for j, s in enumerate(grafico.series):
                color = _SERIES[j % len(_SERIES)]
                if ch.kind == "line":
                    s.graphicalProperties.line.solidFill = color
                    s.graphicalProperties.line.width = 22000
                else:
                    s.graphicalProperties.solidFill = color
                    s.graphicalProperties.line.solidFill = color
            if len(ch.series) == 1:
                grafico.legend = None
            wg.add_chart(grafico, f"{get_column_letter(len(ch.series) + 3)}{fila}")
            fila = max(ultimo + 3, fila + 17)
        wg.column_dimensions["A"].width = 26

    for sec in spec.sections:
        wsec = wb.create_sheet(title=_sheet_title(sec.title, used))
        _header(wsec, 1, [c.header for c in sec.columns])
        for ri, fila_datos in enumerate(sec.rows, start=2):
            for ci, c in enumerate(sec.columns, start=1):
                celda = wsec.cell(row=ri, column=ci, value=_xlsx_safe(fila_datos.get(c.key)))
                if c.money:
                    celda.number_format = _MONEY_FMT
        n = len(sec.rows)
        if sec.total_row:
            for ci, c in enumerate(sec.columns, start=1):
                v = sec.total_row.get(c.key, "")
                celda = wsec.cell(row=n + 2, column=ci, value=_xlsx_safe(v))
                celda.font = Font(bold=True)
                if c.money and v != "":
                    celda.number_format = _MONEY_FMT
        if n:
            for ci, c in enumerate(sec.columns, start=1):
                rango = f"{get_column_letter(ci)}2:{get_column_letter(ci)}{n + 1}"
                alerta = PatternFill("solid", fgColor="FEE2E2")
                if c.alert_below is not None:
                    wsec.conditional_formatting.add(rango, CellIsRule(operator="lessThan", formula=[str(c.alert_below)], fill=alerta, font=Font(color=_BAD)))
                if c.alert_above is not None:
                    wsec.conditional_formatting.add(rango, CellIsRule(operator="greaterThan", formula=[str(c.alert_above)], fill=alerta, font=Font(color=_BAD)))
                if c.data_bar:
                    wsec.conditional_formatting.add(rango, DataBarRule(start_type="min", end_type="max", color=_BRAND))
        for ci, c in enumerate(sec.columns, start=1):
            width = max([len(c.header)] + [len(str(x.get(c.key, ""))) for x in sec.rows]) if sec.rows else len(c.header)
            wsec.column_dimensions[get_column_letter(ci)].width = min(width + 4, 48)
        wsec.freeze_panes = "A2"
        if n:
            wsec.auto_filter.ref = f"A1:{get_column_letter(len(sec.columns))}{n + 1}"
        if sec.note:
            wsec.cell(row=n + (4 if sec.total_row else 3), column=1, value=sec.note).font = Font(italic=True, color=_MUTED)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Specs por pestaña ───

def _ddmm(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%d/%m")


async def spec_ingresos(db: AsyncSession, from_date, to_date) -> ReportSpec:
    s = await rd.summary(db, from_date, to_date)
    prev_from, prev_to = previous_range(from_date, to_date)
    p = await rd.summary(db, prev_from, prev_to)
    kinds = await rd.revenue_by_kind_and_method(db, from_date, to_date)
    plans = await rd.revenue_by_plan(db, from_date, to_date)
    daily = await rd.revenue_daily(db, from_date, to_date)
    expenses = await rd.expenses_by_category(db, from_date, to_date)
    payments = await rd.payments_list(db, from_date, to_date)
    truncated_note = f"Se muestran los primeros {rd.PAYMENTS_LIST_LIMIT} pagos del período." if len(payments) >= rd.PAYMENTS_LIST_LIMIT else None
    return ReportSpec(
        title="Ingresos", subtitle=_period_label(s["from"], s["to"]),
        comparison_label=_period_label(p["from"], p["to"]),
        tiles=[
            _tile("Ingresos confirmados", s["revenue"], p["revenue"], money=True),
            _tile("Egresos", s["expenses"], p["expenses"], money=True, good_when_up=False),
            _tile("Utilidad", s["profit"], p["profit"], money=True),
            _tile("Pagos recibidos", s["payments_count"], p["payments_count"]),
            _tile("Pendiente por confirmar", s["pending"], p["pending"], money=True, good_when_up=False),
            _tile("Devoluciones", s["refunded"], p["refunded"], money=True, good_when_up=False),
            _tile("Socios nuevos", s["new_members"], p["new_members"]),
        ],
        charts=[
            Chart("bar", "Ingresos por día", [_ddmm(d["date"]) for d in daily],
                  [("Membresías", [d["membership"] for d in daily]), ("Pases", [d["day_pass"] for d in daily]),
                   ("Productos", [d["product"] for d in daily])], money=True),
            Chart("hbar", "Ingresos por método de pago", [m["label"] for m in kinds["by_method"]],
                  [("Monto", [m["amount"] for m in kinds["by_method"]])], money=True),
        ],
        sections=[
            Section("Por fuente de ingreso", [Column("label", "Fuente"), Column("count", "Pagos", align_right=True, width_mm=25), Column("amount", "Monto", money=True, width_mm=40, data_bar=True)],
                    kinds["by_kind"], {"label": "TOTAL", "amount": kinds["total"]}),
            Section("Por método de pago", [Column("label", "Método"), Column("count", "Pagos", align_right=True, width_mm=25), Column("amount", "Monto", money=True, width_mm=40, data_bar=True)],
                    kinds["by_method"], {"label": "TOTAL", "amount": kinds["total"]}),
            Section("Por plan", [Column("plan_name", "Plan"), Column("count", "Cobros", align_right=True, width_mm=25), Column("amount", "Monto", money=True, width_mm=40, data_bar=True)], plans),
            Section("Egresos por categoría", [Column("label", "Categoría"), Column("count", "Registros", align_right=True, width_mm=25), Column("amount", "Monto", money=True, width_mm=40)],
                    expenses, {"label": "TOTAL", "amount": s["expenses"]} if expenses else None),
            Section("Ingresos por día", [Column("date", "Fecha", width_mm=30), Column("membership", "Membresías", money=True), Column("day_pass", "Pases", money=True),
                                        Column("product", "Productos", money=True), Column("amount", "Total", money=True, data_bar=True)],
                    [{**d, "date": date.fromisoformat(d["date"]).strftime("%d/%m/%Y")} for d in daily if d["amount"] > 0], {"date": "TOTAL", "amount": s["revenue"]}),
            Section("Detalle de pagos", [Column("date", "Fecha", width_mm=28), Column("member", "Socio"), Column("concept", "Concepto"), Column("method", "Método", width_mm=24),
                                         Column("status", "Estado", width_mm=20), Column("amount", "Monto", money=True, width_mm=26)], payments, note=truncated_note),
        ],
    )


async def spec_ventas(db: AsyncSession, from_date, to_date) -> ReportSpec:
    from app.api.v1.reports import _product_sales_data
    data = await _product_sales_data(db, from_date, to_date)
    prev_from, prev_to = previous_range(from_date, to_date)
    prev = await _product_sales_data(db, prev_from, prev_to)
    stock = await rd.stockout(db)
    top = sorted(data["products"], key=lambda x: x["revenue"], reverse=True)[:10]
    return ReportSpec(
        title="Ventas de productos", subtitle=_period_label(data["from"], data["to"]),
        comparison_label=_period_label(prev["from"], prev["to"]),
        tiles=[
            _tile("Unidades vendidas", data["total_units"], prev["total_units"]),
            _tile("Ingreso neto", data["total_revenue"], prev["total_revenue"], money=True),
            _tile("Costo", data["total_cost"], prev["total_cost"], money=True, good_when_up=False),
            _tile("Margen", data["total_margin"], prev["total_margin"], money=True),
        ],
        charts=[Chart("hbar", "Productos que más venden", [x["name"] for x in top], [("Ingreso neto", [x["revenue"] for x in top])], money=True)],
        sections=[
            Section("Ventas por producto", [Column("name", "Producto"), Column("units", "Unid.", align_right=True, width_mm=18), Column("gross", "Bruto", money=True, width_mm=26),
                                            Column("discount", "Desc.", money=True, width_mm=22), Column("revenue", "Neto", money=True, width_mm=26, data_bar=True),
                                            Column("cost", "Costo", money=True, width_mm=26), Column("margin", "Margen", money=True, width_mm=26, alert_below=0)],
                    data["products"], {"name": "TOTAL", "units": data["total_units"], "revenue": data["total_revenue"], "cost": data["total_cost"], "margin": data["total_margin"]}),
            Section("Inventario actual y proyección de agotamiento (ritmo de 30 días)",
                    [Column("name", "Producto"), Column("stock", "Stock", align_right=True, width_mm=20), Column("units_sold_30d", "Vend. 30d", align_right=True, width_mm=22),
                     Column("daily_velocity", "Por día", align_right=True, width_mm=20), Column("days_left", "Días restantes", align_right=True, width_mm=28, alert_below=8), Column("runs_out_on", "Se agota", width_mm=26)],
                    stock),
        ],
    )


async def spec_pases(db: AsyncSession, from_date, to_date) -> ReportSpec:
    from app.api.v1.reports import _day_passes_data
    data = await _day_passes_data(db, from_date, to_date)
    prev_from, prev_to = previous_range(from_date, to_date)
    prev = await _day_passes_data(db, prev_from, prev_to)
    return ReportSpec(
        title="Pases del día", subtitle=_period_label(data["from"], data["to"]),
        comparison_label=_period_label(prev["from"], prev["to"]),
        tiles=[
            _tile("Pases vendidos", data["total_count"], prev["total_count"]),
            _tile("Ingresos", data["total_amount"], prev["total_amount"], money=True),
            _tile("Personas distintas", data["unique_people"], prev["unique_people"]),
            _tile("Se hicieron socios", data["converted_to_member"], prev["converted_to_member"]),
        ],
        charts=[Chart("bar", "Pases vendidos por día", [_ddmm(d["date"]) for d in data["daily"]], [("Pases", [d["count"] for d in data["daily"]])])],
        sections=[
            Section("Por tipo de pase", [Column("name", "Pase"), Column("includes_trainer", "Entrenador", width_mm=26), Column("count", "Cantidad", align_right=True, width_mm=24), Column("amount", "Ingresos", money=True, width_mm=36, data_bar=True)],
                    [{**t, "includes_trainer": "Sí" if t["includes_trainer"] else "No"} for t in data["by_type"]],
                    {"name": "TOTAL", "count": data["total_count"], "amount": data["total_amount"]}),
            Section("Por día", [Column("date", "Fecha", width_mm=30), Column("count", "Pases", align_right=True, width_mm=24), Column("amount", "Ingresos", money=True, width_mm=36)],
                    [{**d, "date": date.fromisoformat(d["date"]).strftime("%d/%m/%Y")} for d in data["daily"]]),
        ],
    )


async def spec_socios(db: AsyncSession, from_date, to_date) -> ReportSpec:
    from app.services.exports import _expiring_rows, _inactive_rows, _roster_rows
    o = await rd.members_overview(db, from_date, to_date)
    prev_from, prev_to = previous_range(from_date, to_date)
    po = await rd.members_overview(db, prev_from, prev_to)
    st = o["by_membership_status"]
    labels = {"active": "Con membresía activa", "scheduled": "Programada", "frozen": "Congelada", "in_grace": "En gracia", "expired": "Vencida", "no_membership": "Sin membresía"}
    inactivos = await _inactive_rows(db)
    return ReportSpec(
        title="Socios", subtitle=_period_label(o["from"], o["to"]),
        comparison_label=_period_label(po["from"], po["to"]),
        tiles=[
            Tile("Socios registrados", str(o["total"]), raw=o["total"]),
            Tile("Cuentas activas", str(o["active_accounts"]), raw=o["active_accounts"]),
            _tile("Nuevos en el período", o["new_in_period"], po["new_in_period"]),
            _tile("Visitantes (pase)", o["visitors"], po["visitors"]),
            Tile("Sin venir hace 14+ días", str(len(inactivos)), raw=len(inactivos)),
        ],
        charts=[Chart("hbar", "Estado de membresía", [labels.get(k, k) for k in st], [("Socios", list(st.values()))])],
        sections=[
            Section("Estado de membresía (socios activos)", [Column("label", "Estado"), Column("count", "Socios", align_right=True, width_mm=30, data_bar=True)],
                    [{"label": labels.get(k, k), "count": v} for k, v in st.items()]),
            Section("Nuevos en el período", [Column("full_name", "Socio"), Column("phone", "Teléfono", width_mm=34), Column("created_at", "Registro", width_mm=28)], o["new_list"]),
            Section("Membresías por vencer (30 días)", [Column("member_name", "Socio"), Column("plan_name", "Plan"), Column("end_date", "Vence", width_mm=28), Column("days_remaining", "Días", align_right=True, width_mm=18, alert_below=8)],
                    await _expiring_rows(db)),
            Section("Sin venir hace 14+ días", [Column("full_name", "Socio"), Column("phone", "Teléfono", width_mm=34), Column("last_checkin", "Última visita", width_mm=30), Column("days_inactive", "Días", align_right=True, width_mm=18, alert_above=30)],
                    inactivos),
            Section("Listado de socios", [Column("full_name", "Nombre"), Column("id_number", "Documento", width_mm=30), Column("phone", "Teléfono", width_mm=32), Column("email", "Email"), Column("status", "Estado", width_mm=20)],
                    await _roster_rows(db)),
        ],
    )


_SOURCE_LABELS = {
    "reception": "Recepción", "kiosk": "Kiosco", "qr": "Código QR", "turnstile": "Torniquete",
    "fingerprint": "Huella", "face": "Rostro", "app": "App", "class": "Clase",
}


def _busiest(peak: dict) -> str:
    if peak["busiest_hour"] is None:
        return "—"
    return f"{rd.eng.WEEKDAYS_ES[peak['busiest_weekday']]} · {peak['busiest_hour']:02d}:00"


async def spec_asistencia(db: AsyncSession, from_date, to_date) -> ReportSpec:
    a = await rd.attendance(db, from_date, to_date)
    prev_from, prev_to = previous_range(from_date, to_date)
    pa = await rd.attendance(db, prev_from, prev_to)
    peak = a["peak"]
    horas = [h for h in peak["by_hour"] if h["count"] > 0]
    return ReportSpec(
        title="Asistencia", subtitle=_period_label(a["from"], a["to"]),
        comparison_label=_period_label(pa["from"], pa["to"]),
        tiles=[
            _tile("Ingresos al gym", a["total"], pa["total"]),
            _tile("Personas distintas", a["unique_members"], pa["unique_members"]),
            _tile("Accesos denegados", a["denied"], pa["denied"], good_when_up=False),
            Tile("Momento pico", _busiest(peak)),
        ],
        charts=[
            Chart("bar", "Ingresos por hora", [f"{h['hour']:02d}h" for h in horas], [("Ingresos", [h["count"] for h in horas])]),
            Chart("bar", "Ingresos por día de la semana", [d["label"] for d in peak["by_weekday"]], [("Ingresos", [d["count"] for d in peak["by_weekday"]])]),
        ],
        sections=[
            Section("Por día", [Column("date", "Fecha", width_mm=30), Column("count", "Ingresos", align_right=True, width_mm=30, data_bar=True)],
                    [{**d, "date": date.fromisoformat(d["date"]).strftime("%d/%m/%Y")} for d in a["daily"] if d["count"] > 0],
                    {"date": "TOTAL", "count": a["total"]}, empty_text="Nadie entró al gimnasio en el período."),
            Section("Por día de la semana", [Column("label", "Día"), Column("count", "Ingresos", align_right=True, width_mm=30, data_bar=True)], peak["by_weekday"]),
            Section("Por hora", [Column("hour", "Hora", align_right=True, width_mm=20), Column("count", "Ingresos", align_right=True, width_mm=30, data_bar=True)],
                    [{"hour": f"{h['hour']:02d}:00", "count": h["count"]} for h in horas]),
            Section("Por origen", [Column("source", "Origen"), Column("count", "Ingresos", align_right=True, width_mm=30)],
                    [{**s, "source": _SOURCE_LABELS.get(s["source"], s["source"])} for s in a["by_source"]]),
        ],
    )


async def spec_retencion(db: AsyncSession, from_date, to_date) -> ReportSpec:
    from app.api.v1.reports import _retention_data
    ret = await _retention_data(db, 6)
    coh = await rd.cohorts(db, 6)
    risk = await rd.churn_risk(db, 50)
    con_datos = [r for r in ret if r["renewal_rate"] is not None]
    last = con_datos[-1] if con_datos else None
    before = con_datos[-2] if len(con_datos) > 1 else None
    tiles = [
        Tile("Renovación (último mes con datos)", f"{last['renewal_rate']}%" if last else "—",
             previous=f"{before['renewal_rate']}%" if before else None,
             delta_pct=_delta(last["renewal_rate"], before["renewal_rate"]) if last and before else None),
        Tile("Churn", f"{last['churn_rate']}%" if last else "—",
             previous=f"{before['churn_rate']}%" if before else None,
             delta_pct=_delta(last["churn_rate"], before["churn_rate"]) if last and before else None, good_when_up=False),
        Tile("Socios en riesgo alto", str(sum(1 for r in risk if r["level"] == "alto"))),
    ]
    return ReportSpec(
        title="Retención", subtitle="Últimos 6 meses",
        comparison_label="el mes anterior con datos" if before else None,
        tiles=tiles,
        charts=[Chart("line", "Renovación y churn por mes (%)", [r["month"] for r in con_datos],
                      [("Renovación", [float(r["renewal_rate"]) for r in con_datos]), ("Churn", [float(r["churn_rate"] or 0) for r in con_datos])])],
        sections=[
            Section("Renovación mensual", [Column("month", "Mes", width_mm=30), Column("ended", "Vencieron", align_right=True, width_mm=26), Column("renewed", "Renovaron", align_right=True, width_mm=26),
                                           Column("renewal_rate", "Renovación %", align_right=True, width_mm=28), Column("churn_rate", "Churn %", align_right=True, width_mm=24, alert_above=15)], ret),
            Section("Cohortes por mes de ingreso", [Column("month", "Cohorte", width_mm=30), Column("size", "Socios", align_right=True, width_mm=22), Column("active_now", "Activos hoy", align_right=True, width_mm=26),
                                                     Column("active_pct", "Activos %", align_right=True, width_mm=24, data_bar=True), Column("renewed", "Renovaron", align_right=True, width_mm=24), Column("renewed_pct", "Renov. %", align_right=True, width_mm=22)], coh),
            Section("Riesgo de abandono", [Column("full_name", "Socio"), Column("phone", "Teléfono", width_mm=30), Column("score", "Puntaje", align_right=True, width_mm=20, alert_above=59), Column("level", "Nivel", width_mm=18), Column("reasons_txt", "Por qué")],
                    [{**r, "reasons_txt": "; ".join(r["reasons"])} for r in risk]),
        ],
    )


async def spec_inteligencia(db: AsyncSession, from_date, to_date) -> ReportSpec:
    """Exporta lo mismo que muestra la pestaña Inteligencia: proyección de
    ingresos, horas pico, riesgo de abandono y stock por agotarse."""
    fc = await rd.forecast(db)
    a = await rd.attendance(db, from_date, to_date)
    peak = a["peak"]
    risk = await rd.churn_risk(db, 50)
    stock = await rd.stockout(db)
    horas = [h for h in peak["by_hour"] if h["count"] > 0]
    return ReportSpec(
        title="Inteligencia", subtitle=_period_label(a["from"], a["to"]),
        comparison_label="el promedio de los 3 meses previos" if fc.get("avg_previous_months") is not None else None,
        tiles=[
            Tile("Proyección del mes", _cop(fc["projection"]), raw=fc["projection"], money=True,
                 previous=_cop(fc["avg_previous_months"]) if fc.get("avg_previous_months") is not None else None,
                 raw_previous=fc.get("avg_previous_months"), delta_pct=fc.get("vs_previous_pct")),
            Tile("Va del mes", _cop(fc["month_to_date"]), raw=fc["month_to_date"], money=True),
            Tile("Hora pico", _busiest(peak)),
            Tile("Socios en riesgo alto", str(sum(1 for r in risk if r["level"] == "alto"))),
        ],
        charts=[Chart("bar", "Horas pico", [f"{h['hour']:02d}h" for h in horas], [("Ingresos", [h["count"] for h in horas])])],
        sections=[
            Section("Proyección de ingresos del mes",
                    [Column("concepto", "Concepto"), Column("valor", "Valor", align_right=True, width_mm=40)],
                    [{"concepto": "Ingresos en lo que va del mes", "valor": _cop(fc["month_to_date"])},
                     {"concepto": f"Ritmo diario ({fc['elapsed_days']} de {fc['days_in_month']} días)", "valor": _cop(fc["daily_pace"])},
                     {"concepto": "Proyección al cierre", "valor": _cop(fc["projection"])},
                     {"concepto": "Rango probable", "valor": f"{_cop(fc['projection_low'])} a {_cop(fc['projection_high'])}"},
                     {"concepto": "Promedio de los 3 meses previos",
                      "valor": _cop(fc["avg_previous_months"]) if fc.get("avg_previous_months") is not None else "sin historial"}]),
            Section("Horas pico", [Column("hour", "Hora", align_right=True, width_mm=20), Column("count", "Ingresos", align_right=True, width_mm=30, data_bar=True)],
                    [{"hour": f"{h['hour']:02d}:00", "count": h["count"]} for h in horas]),
            Section("Días más movidos", [Column("label", "Día"), Column("count", "Ingresos", align_right=True, width_mm=30, data_bar=True)], peak["by_weekday"]),
            Section("Socios en riesgo de abandono",
                    [Column("full_name", "Socio"), Column("phone", "Teléfono", width_mm=30),
                     Column("score", "Puntaje", align_right=True, width_mm=20, alert_above=59), Column("level", "Nivel", width_mm=18),
                     Column("reasons_txt", "Por qué")],
                    [{**r, "reasons_txt": "; ".join(r["reasons"])} for r in risk]),
            Section("Inventario por agotarse",
                    [Column("name", "Producto"), Column("stock", "Stock", align_right=True, width_mm=20),
                     Column("units_sold_30d", "Vendidos 30d", align_right=True, width_mm=28),
                     Column("days_left_txt", "Días restantes", align_right=True, width_mm=28, alert_below=8),
                     Column("runs_out_on_txt", "Se agota", width_mm=26)],
                    [{**p, "days_left_txt": p["days_left"] if p["days_left"] is not None else "—",
                      "runs_out_on_txt": date.fromisoformat(p["runs_out_on"]).strftime("%d/%m/%Y") if p.get("runs_out_on") else "—"}
                     for p in stock]),
        ],
    )


async def spec_cierre(db: AsyncSession, year: int, month: int) -> ReportSpec:
    start = date(year, month, 1)
    end = (date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)) - timedelta(days=1)
    ing = await spec_ingresos(db, start, end)
    ven = await spec_ventas(db, start, end)
    pas = await spec_pases(db, start, end)
    soc = await spec_socios(db, start, end)
    asi = await spec_asistencia(db, start, end)
    ret = await spec_retencion(db, start, end)
    label = f"{MESES_ES[month - 1].capitalize()} {year}"
    generated = today_bogota().strftime("%d/%m/%Y")
    # Secciones que describen el estado de HOY, no el mes: se marcan para no confundir un cierre pasado
    snapshot_titles = {"Inventario actual y proyección de agotamiento (ritmo de 30 días)", "Membresías por vencer (30 días)",
                       "Sin venir hace 14+ días", "Renovación mensual", "Cohortes por mes de ingreso", "Riesgo de abandono"}
    sections = []
    for sp, prefix in ((ing, "Ingresos"), (ven, "Ventas"), (pas, "Pases"), (soc, "Socios"), (asi, "Asistencia"), (ret, "Retención")):
        for sec in sp.sections:
            if sec.title in ("Detalle de pagos", "Listado de socios"):
                continue
            title = f"{prefix} · {sec.title}"
            if sec.title in snapshot_titles:
                title += f" (estado al {generated})"
            sections.append(Section(title, sec.columns, sec.rows, sec.total_row, sec.note, sec.empty_text))
    por_titulo = {t.label: t for sp in (ing, ven, pas, soc, asi, ret) for t in sp.tiles}
    tiles = [por_titulo[k] for k in (
        "Ingresos confirmados", "Egresos", "Utilidad", "Socios nuevos",
        "Ingresos al gym", "Margen", "Pases vendidos", "Renovación (último mes con datos)",
    ) if k in por_titulo]
    return ReportSpec(
        title="Cierre de mes", subtitle=label, comparison_label=ing.comparison_label,
        tiles=tiles, charts=ing.charts[:1] + asi.charts[:1] + ven.charts,
        sections=sections,
    )


SPECS = {
    "ingresos": spec_ingresos,
    "ventas": spec_ventas,
    "pases": spec_pases,
    "socios": spec_socios,
    "asistencia": spec_asistencia,
    "retencion": spec_retencion,
    "inteligencia": spec_inteligencia,
}
