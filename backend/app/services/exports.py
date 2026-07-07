"""Generación server-side de reportes exportables (Excel .xlsx y PDF).

Una sola fuente de verdad para los datos del export: las mismas consultas que
alimentan /reports, ensambladas aquí en datasets tabulares. El frontend solo
dispara la descarga."""
import io
import logging
from datetime import date, datetime, timedelta
from decimal import Decimal

import httpx
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as PdfImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.timeutil import today_bogota
from app.models.tenant.checkin import CheckIn
from app.models.tenant.expense import Expense
from app.models.tenant.member import Member, MemberStatus
from app.models.tenant.membership import Membership, MembershipStatus
from app.models.tenant.payment import Payment, PaymentStatus

logger = logging.getLogger(__name__)

_VOLT = "D4FF00"  # color de marca para encabezados
_INK = "111111"


async def _is_public_host(url: str) -> bool:
    """Anti-SSRF: el logo_url lo escribe el gym; nunca debe alcanzar IPs internas
    (metadata de EC2, servicios del compose, loopback)."""
    import asyncio
    import ipaddress
    import socket
    from urllib.parse import urlparse

    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        return False
    try:
        infos = await asyncio.to_thread(socket.getaddrinfo, parsed.hostname, None)
    except OSError:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return False
    return True


async def fetch_logo_png(logo_url: str | None) -> bytes | None:
    """Descarga el logo del gym y lo normaliza a PNG (fallback silencioso: sin logo)."""
    if not logo_url:
        return None
    try:
        from app.core.config import settings as app_settings

        # El propio bucket S3/MinIO del sistema es confiable aunque sea interno (dev).
        # La URL pública (localhost:9000) no resuelve dentro del contenedor: para la
        # descarga se reescribe al endpoint interno (minio:9000).
        own_bucket = bool(app_settings.s3_endpoint_url) and logo_url.startswith(app_settings.s3_endpoint_url)
        public_base = getattr(app_settings, "s3_public_url", "") or ""
        if public_base and app_settings.s3_endpoint_url and logo_url.startswith(public_base):
            logo_url = app_settings.s3_endpoint_url.rstrip("/") + logo_url[len(public_base.rstrip("/")):]
            own_bucket = True
        if not own_bucket and not await _is_public_host(logo_url):
            return None
        async with httpx.AsyncClient(timeout=5.0, follow_redirects=True) as client:
            resp = await client.get(logo_url)
            resp.raise_for_status()
            raw = resp.content[: 3 * 1024 * 1024]
        img = PILImage.open(io.BytesIO(raw))
        img.thumbnail((480, 480))
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGBA")
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()
    except Exception:
        logger.info("Logo del gym no disponible para el export (%s)", logo_url)
        return None


def _month_bounds(year: int, month: int) -> tuple[datetime, datetime]:
    start = date(year, month, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    return datetime.combine(start, datetime.min.time()), datetime.combine(end, datetime.min.time())


def _months_back(today: date, i: int) -> tuple[int, int]:
    m, y = today.month - i, today.year
    while m <= 0:
        m += 12
        y -= 1
    return y, m


async def _profit_rows(db: AsyncSession, months: int = 12) -> list[dict]:
    today = today_bogota()
    rows = []
    for i in range(months - 1, -1, -1):
        y, m = _months_back(today, i)
        start_dt, end_dt = _month_bounds(y, m)
        income = (await db.execute(
            select(func.coalesce(func.sum(Payment.amount), Decimal("0"))).where(
                Payment.created_at >= start_dt, Payment.created_at < end_dt,
                Payment.status == PaymentStatus.paid,
            )
        )).scalar() or Decimal("0")
        expenses = (await db.execute(
            select(func.coalesce(func.sum(Expense.amount), Decimal("0"))).where(
                Expense.created_at >= start_dt, Expense.created_at < end_dt,
            )
        )).scalar() or Decimal("0")
        rows.append({
            "month": date(y, m, 1).strftime("%b %Y"),
            "income": float(income),
            "expenses": float(expenses),
            "profit": float(income - expenses),
        })
    return rows


async def _expiring_rows(db: AsyncSession, days: int = 30) -> list[dict]:
    today = today_bogota()
    cutoff = today + timedelta(days=days)
    result = await db.execute(
        select(Membership, Member.full_name)
        .join(Member, Member.id == Membership.member_id)
        .where(
            Membership.end_date <= cutoff,
            Membership.end_date >= today,
            Membership.status == MembershipStatus.active.value,
        )
        .order_by(Membership.end_date)
    )
    return [
        {
            "member_name": full_name,
            "plan_name": m.plan_name,
            "end_date": m.end_date.isoformat(),
            "days_remaining": (m.end_date - today).days,
        }
        for m, full_name in result.all()
    ]


async def _inactive_rows(db: AsyncSession, days: int = 14) -> list[dict]:
    today = today_bogota()
    cutoff = datetime.combine(today - timedelta(days=days), datetime.min.time())
    recent = (
        select(CheckIn.member_id).where(CheckIn.created_at >= cutoff).distinct().scalar_subquery()
    )
    members = (await db.execute(
        select(Member)
        .where(Member.status == MemberStatus.active, Member.id.not_in(recent))
        .order_by(Member.full_name)
    )).scalars().all()
    ids = [m.id for m in members]
    last_map: dict = {}
    if ids:
        lr = await db.execute(
            select(CheckIn.member_id, func.max(CheckIn.created_at))
            .where(CheckIn.member_id.in_(ids)).group_by(CheckIn.member_id)
        )
        last_map = {row[0]: row[1] for row in lr.all()}
    return [
        {
            "full_name": m.full_name,
            "phone": m.phone or "",
            "email": m.email or "",
            "last_checkin": last_map[m.id].date().isoformat() if last_map.get(m.id) else "Nunca",
            "days_inactive": (today - last_map[m.id].date()).days if last_map.get(m.id) else "",
        }
        for m in members
    ]


async def _roster_rows(db: AsyncSession) -> list[dict]:
    members = (await db.execute(
        select(Member).order_by(Member.full_name)
    )).scalars().all()
    return [
        {
            "full_name": m.full_name,
            "id_number": m.id_number or "",
            "phone": m.phone or "",
            "email": m.email or "",
            "status": "Activo" if m.status == MemberStatus.active else "Inactivo",
        }
        for m in members
    ]


# ─────────────────────────── Excel ───────────────────────────

def _write_sheet(ws, title: str, columns: list[tuple[str, str]], rows: list[dict]) -> None:
    """columns: lista de (clave, encabezado). Aplica estilo de marca al header."""
    ws.title = title
    header_font = Font(bold=True, color="111111")
    header_fill = PatternFill("solid", fgColor=_VOLT)
    for col_idx, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key))
    # Ancho aproximado por contenido
    for col_idx, (key, header) in enumerate(columns, start=1):
        width = max(len(header), *(len(str(r.get(key, ""))) for r in rows)) if rows else len(header)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 4, 45)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


_MONEY_FMT = '"$"#,##0'


def _apply_money_format(ws, money_cols: list[int], n_rows: int) -> None:
    for col_idx in money_cols:
        for row_idx in range(2, n_rows + 2):
            ws.cell(row=row_idx, column=col_idx).number_format = _MONEY_FMT


def _cover_sheet(wb: Workbook, gym_name: str, logo_png: bytes | None) -> None:
    ws = wb.active
    ws.title = "Portada"
    ws.sheet_view.showGridLines = False
    row = 2
    if logo_png:
        img = XlsxImage(io.BytesIO(logo_png))
        scale = min(1.0, 160 / max(img.width, 1))
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
        ws.add_image(img, "B2")
        row = 12
    ws.cell(row=row, column=2, value=gym_name).font = Font(bold=True, size=22, color=_INK)
    ws.cell(row=row + 1, column=2, value="Reporte operativo y financiero").font = Font(size=12, color="666666")
    ws.cell(row=row + 2, column=2, value=f"Generado el {today_bogota().strftime('%d/%m/%Y')} · FitConnect").font = Font(size=10, color="999999")
    ws.column_dimensions["B"].width = 50


async def build_reports_xlsx(db: AsyncSession, gym_name: str = "Gimnasio", logo_png: bytes | None = None) -> bytes:
    wb = Workbook()
    _cover_sheet(wb, gym_name, logo_png)
    profit_rows = await _profit_rows(db)
    ws_profit = wb.create_sheet()
    _write_sheet(
        ws_profit, "Utilidad",
        [("month", "Mes"), ("income", "Ingresos"), ("expenses", "Egresos"), ("profit", "Utilidad")],
        profit_rows,
    )
    _apply_money_format(ws_profit, [2, 3, 4], len(profit_rows))
    _write_sheet(
        wb.create_sheet(), "Por vencer",
        [("member_name", "Socio"), ("plan_name", "Plan"), ("end_date", "Vence"), ("days_remaining", "Días")],
        await _expiring_rows(db),
    )
    _write_sheet(
        wb.create_sheet(), "Inactivos",
        [("full_name", "Socio"), ("phone", "Teléfono"), ("email", "Email"),
         ("last_checkin", "Última visita"), ("days_inactive", "Días sin venir")],
        await _inactive_rows(db),
    )
    _write_sheet(
        wb.create_sheet(), "Socios",
        [("full_name", "Nombre"), ("id_number", "Documento"), ("phone", "Teléfono"),
         ("email", "Email"), ("status", "Estado")],
        await _roster_rows(db),
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────── PDF ───────────────────────────

def _cop(v: float) -> str:
    return "$" + f"{v:,.0f}".replace(",", ".")


_TABLE_STYLE = TableStyle([
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + _VOLT)),
    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#" + _INK)),
    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
    ("FONTSIZE", (0, 0), (-1, -1), 9),
    ("TOPPADDING", (0, 0), (-1, -1), 5),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
])


def _section_title(text: str, styles) -> Paragraph:
    return Paragraph(
        text,
        ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                       textColor=colors.HexColor("#" + _INK), spaceBefore=10, spaceAfter=4),
    )


async def _summary_counts(db: AsyncSession) -> dict:
    total = (await db.execute(select(func.count(Member.id)))).scalar() or 0
    active = (await db.execute(
        select(func.count(Member.id)).where(Member.status == MemberStatus.active)
    )).scalar() or 0
    today = today_bogota()
    month_start = datetime.combine(today.replace(day=1), datetime.min.time())
    month_income = (await db.execute(
        select(func.coalesce(func.sum(Payment.amount), Decimal("0"))).where(
            Payment.created_at >= month_start, Payment.status == PaymentStatus.paid,
        )
    )).scalar() or Decimal("0")
    return {"total": total, "active": active, "month_income": float(month_income)}


async def build_report_pdf(db: AsyncSession, gym_name: str = "Gimnasio", logo_png: bytes | None = None) -> bytes:
    rows = await _profit_rows(db)
    expiring = await _expiring_rows(db)
    inactive = await _inactive_rows(db)
    summary = await _summary_counts(db)
    generated = today_bogota().strftime("%d/%m/%Y")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=22 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=20, spaceAfter=2, alignment=0)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], textColor=colors.grey, fontSize=9)

    def _page_decor(canvas, _doc):
        canvas.saveState()
        # Banda de marca superior
        canvas.setFillColor(colors.HexColor("#" + _VOLT))
        canvas.rect(0, A4[1] - 6 * mm, A4[0], 6 * mm, stroke=0, fill=1)
        # Pie: gym + página
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(18 * mm, 12 * mm, f"{gym_name} · Reporte FitConnect · {generated}")
        canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, f"Página {canvas.getPageNumber()}")
        canvas.restoreState()

    elements = []
    if logo_png:
        reader = ImageReader(io.BytesIO(logo_png))
        iw, ih = reader.getSize()
        w = 28 * mm
        elements.append(PdfImage(io.BytesIO(logo_png), width=w, height=w * ih / iw))
        elements.append(Spacer(1, 4 * mm))
    elements += [
        Paragraph(gym_name, title_style),
        Paragraph(f"Reporte operativo y financiero · generado {generated}", sub_style),
        Spacer(1, 6 * mm),
    ]

    # Resumen
    elements.append(_section_title("Resumen", styles))
    summary_table = Table(
        [["Socios registrados", "Socios activos", "Ingresos del mes"],
         [str(summary["total"]), str(summary["active"]), _cop(summary["month_income"])]],
        colWidths=[56 * mm, 56 * mm, 56 * mm],
    )
    summary_table.setStyle(_TABLE_STYLE)
    elements.append(summary_table)

    # Estado de resultados
    elements.append(_section_title("Estado de resultados (últimos 12 meses)", styles))
    data = [["Mes", "Ingresos", "Egresos", "Utilidad"]]
    tot_i = tot_e = tot_p = 0.0
    for r in rows:
        data.append([r["month"], _cop(r["income"]), _cop(r["expenses"]), _cop(r["profit"])])
        tot_i += r["income"]
        tot_e += r["expenses"]
        tot_p += r["profit"]
    data.append(["TOTAL", _cop(tot_i), _cop(tot_e), _cop(tot_p)])
    table = Table(data, colWidths=[42 * mm, 42 * mm, 42 * mm, 42 * mm])
    table.setStyle(_TABLE_STYLE)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    elements.append(table)

    # Membresías por vencer
    elements.append(_section_title("Membresías por vencer (30 días)", styles))
    if expiring:
        exp_data = [["Socio", "Plan", "Vence", "Días"]] + [
            [e["member_name"], e["plan_name"], e["end_date"], str(e["days_remaining"])]
            for e in expiring[:40]
        ]
        exp_table = Table(exp_data, colWidths=[62 * mm, 50 * mm, 34 * mm, 22 * mm])
        exp_table.setStyle(_TABLE_STYLE)
        elements.append(exp_table)
    else:
        elements.append(Paragraph("Ninguna membresía vence en los próximos 30 días.", sub_style))

    # Socios inactivos
    elements.append(_section_title("Socios sin venir (14+ días)", styles))
    if inactive:
        ina_data = [["Socio", "Teléfono", "Última visita", "Días"]] + [
            [i["full_name"], i["phone"], i["last_checkin"], str(i["days_inactive"])]
            for i in inactive[:40]
        ]
        ina_table = Table(ina_data, colWidths=[62 * mm, 40 * mm, 40 * mm, 26 * mm])
        ina_table.setStyle(_TABLE_STYLE)
        elements.append(ina_table)
    else:
        elements.append(Paragraph("Todos los socios activos han visitado recientemente.", sub_style))

    doc.build(elements, onFirstPage=_page_decor, onLaterPages=_page_decor)
    return buf.getvalue()
