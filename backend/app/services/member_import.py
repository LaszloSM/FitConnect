"""Importación de socios desde el software anterior del gimnasio.

Un gimnasio que llega a LSMFit trae su base en la exportación de otro
programa: encabezados propios ("Cédula", "Sexo", "Fecha Nac", "RH"...), fechas
dd/mm/aaaa, cédulas con puntos, nombres en mayúsculas, correos compartidos por
familias. La importación tiene que absorber todo eso sin crear duplicados ni
pisar lo que el gym ya registró.

Flujo en dos pasos sobre el mismo archivo:

1. `preview`: lee, detecta columnas, valida y clasifica cada persona (nueva,
   se completa, sin cambios, con error) sin escribir nada.
2. `apply`: repite exactamente el mismo análisis con el mapeo que confirmó el
   usuario y lo aplica como un lote (`ImportBatch`) que se puede deshacer.

Reglas para que no haya conflictos:
- La llave de una persona es su documento normalizado (sin puntos ni espacios).
- Si ya existe, por defecto solo se llenan los campos vacíos; lo distinto se
  reporta y solo se sobrescribe si el usuario lo elige (`on_conflict="file"`).
- Volver a subir el mismo archivo no duplica socios ni membresías.
- Un correo que ya usa otra persona (familias) no bloquea la fila: entra sin
  correo y queda marcado.
"""
from __future__ import annotations

import asyncio
import csv
import io
import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.timeutil import today_bogota, utcnow_naive
from app.models.tenant.checkin import CheckIn
from app.models.tenant.class_enrollment import ClassEnrollment
from app.models.tenant.day_pass import DayPass
from app.models.tenant.import_batch import ImportBatch
from app.models.tenant.instructor import Instructor
from app.models.tenant.member import Gender, IdType, Member, MemberType, TrainingGoal
from app.models.tenant.member_plan import MemberPlan
from app.models.tenant.membership import Membership, MembershipStatus
from app.models.tenant.membership_freeze import MembershipFreeze
from app.models.tenant.payment import Payment
from app.models.tenant.routine import MemberRoutine
from app.models.tenant.session_booking import SessionBooking
from app.schemas.member import MemberBase
from app.services.exports import _BRAND_TINT

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
# Filas que viajan en la vista previa; los conteos siempre cubren el archivo entero.
PREVIEW_ROWS = 2_000
IMPORTED_PLAN_NAME = "Importado del sistema anterior"


class ImportFileError(ValueError):
    """El archivo no se puede leer o excede los límites."""


# ─── Catálogo de campos ──────────────────────────────────────────────────────

@dataclass(frozen=True)
class FieldSpec:
    key: str
    label: str
    aliases: tuple[str, ...]
    #: Varias columnas pueden alimentar el mismo campo (Nombres + Apellidos,
    #: Grupo + RH); sus valores se unen en el orden de las columnas.
    joinable: bool = False


FIELDS: tuple[FieldSpec, ...] = (
    FieldSpec("id_number", "Cédula / documento", (
        "cedula", "cedula de ciudadania", "cc", "c c", "documento", "numero de documento",
        "no documento", "nro documento", "num documento", "numero documento", "documento de identidad",
        "identificacion", "numero de identificacion", "no identificacion", "dni", "nro identificacion",
    )),
    FieldSpec("id_type", "Tipo de documento", (
        "tipo de documento", "tipo documento", "tipo doc", "tipo id", "tipo de identificacion",
        "tipo identificacion",
    )),
    FieldSpec("full_name", "Nombre", (
        "nombre", "nombres", "nombre completo", "apellido", "apellidos", "nombres y apellidos",
        "nombre y apellido", "nombre y apellidos", "cliente", "socio", "miembro", "primer nombre",
        "segundo nombre", "primer apellido", "segundo apellido",
    ), joinable=True),
    FieldSpec("gender", "Sexo", ("sexo", "genero", "sexo genero")),
    FieldSpec("birth_date", "Fecha de nacimiento", (
        "fecha nac", "fecha nacimiento", "fecha de nacimiento", "nacimiento", "f nacimiento",
        "fec nacimiento", "fecha nac aaaa mm dd", "fecha nac dd mm aaaa",
    )),
    FieldSpec("city", "Ciudad", ("ciudad", "municipio", "ciudad de residencia")),
    FieldSpec("email", "Email", ("email", "e mail", "correo", "correo electronico", "mail")),
    FieldSpec("phone", "Teléfono", (
        "telefono", "celular", "movil", "tel", "whatsapp", "telefono celular", "numero de celular",
        "numero celular", "telefono movil", "cel",
    )),
    FieldSpec("eps", "EPS", ("eps", "entidad de salud", "aseguradora", "salud", "entidad promotora de salud")),
    FieldSpec("blood_type", "RH / tipo de sangre", (
        "rh", "tipo de sangre", "grupo sanguineo", "gs rh", "gs", "grupo sanguineo y rh", "sangre",
        "factor rh", "tipo sangre", "grupo y rh",
    ), joinable=True),
    FieldSpec("expectation", "Expectativa / objetivo", (
        "expectativa", "expectativas", "objetivo", "objetivos", "meta", "metas",
    )),
    FieldSpec("address", "Dirección", ("direccion", "direccion de residencia", "domicilio", "dir")),
    FieldSpec("access_by_document", "Ingreso por cédula", (
        "ingreso por dni", "ingreso por cedula", "ingreso con cedula", "ingreso por documento",
        "acceso por documento", "acceso por cedula", "entrada por cedula", "ingreso dni",
    )),
    FieldSpec("trainer", "Entrenador", (
        "entrenador", "instructor", "coach", "profesor", "entrenador asignado", "entrenador personal",
    )),
    FieldSpec("emergency_contact_name", "Contacto de emergencia", (
        "contacto de emergencia", "contacto emergencia", "nombre contacto emergencia", "acudiente",
        "nombre de emergencia", "en caso de emergencia",
    )),
    FieldSpec("emergency_contact_phone", "Teléfono de emergencia", (
        "telefono de emergencia", "telefono emergencia", "tel emergencia", "celular emergencia",
        "telefono contacto emergencia", "telefono acudiente",
    )),
    FieldSpec("country", "País", ("pais",)),
    FieldSpec("notes", "Notas", ("notas", "observaciones", "observacion", "comentarios", "nota"), joinable=True),
    FieldSpec("plan", "Plan", (
        "plan", "membresia", "tipo de plan", "plan actual", "tipo de membresia", "servicio", "tiquetera",
        "ticketera",
    )),
    FieldSpec("membership_start", "Fecha inicio del plan", (
        "fecha inicio", "inicio", "fecha de inicio", "desde", "fecha inicio plan", "inicio plan",
    )),
    FieldSpec("membership_end", "Fecha vencimiento del plan", (
        "fecha vencimiento", "vencimiento", "vence", "fecha fin", "fecha final", "hasta",
        "fecha de vencimiento", "fecha de corte", "fecha vence", "fin plan",
    )),
    FieldSpec("entries_remaining", "Entradas restantes", (
        "entradas restantes", "entradas", "ingresos restantes", "saldo entradas", "tiquetes",
        "tickets", "clases restantes", "sesiones restantes", "entradas disponibles",
    )),
)
FIELD_BY_KEY = {f.key: f for f in FIELDS}

# Campos de la ficha del socio que la importación puede llenar o comparar.
MEMBER_FIELDS = (
    "full_name", "id_type", "id_number", "gender", "birth_date", "city", "email", "phone", "eps",
    "blood_type", "expectation", "training_goal", "address", "access_by_document", "trainer_id",
    "emergency_contact_name", "emergency_contact_phone", "country", "notes",
)


def _norm(text: Any) -> str:
    """Minúsculas, sin tildes ni signos: "Fecha Nac.*" → "fecha nac"."""
    s = unicodedata.normalize("NFKD", str(text or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return s.strip()


_ALIAS: dict[str, str] = {}
for _f in FIELDS:
    _ALIAS[_norm(_f.label)] = _f.key
    _ALIAS[_f.key.replace("_", " ")] = _f.key
    for _a in _f.aliases:
        _ALIAS[_norm(_a)] = _f.key
# Encabezados de la plantilla anterior de LSMFit (siguen funcionando).
_ALIAS.update({
    "numero de documento": "id_number", "genero": "gender", "correo": "email",
    "tipo de sangre": "blood_type", "objetivo": "expectation", "pais 2 letras": "country",
})


def detect_field(header: Any) -> Optional[str]:
    return _ALIAS.get(_norm(header))


# ─── Lectura del archivo ─────────────────────────────────────────────────────

@dataclass
class Sheet:
    headers: list[str]
    rows: list[tuple[int, list[Any]]]  # (número de fila en el archivo, celdas)
    sheet_name: Optional[str] = None


def _decode_csv(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("No se pudo leer el texto del archivo")


def _read_csv(data: bytes) -> list[list[Any]]:
    text = _decode_csv(data)
    sample = text[:8192]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(data: bytes) -> tuple[list[list[Any]], str]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("No se pudo abrir el Excel. Guárdalo como .xlsx, .xls o .csv y vuelve a intentar.")
    ws = wb["Miembros"] if "Miembros" in wb.sheetnames else wb.worksheets[0]
    rows: list[list[Any]] = []
    for raw in ws.iter_rows(values_only=True):
        rows.append(list(raw))
        if len(rows) > MAX_ROWS + 20:
            break
    return rows, ws.title


def _read_xls(data: bytes) -> tuple[list[list[Any]], str]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception:
        raise ImportFileError("No se pudo abrir el Excel (.xls). Prueba guardarlo como .xlsx o .csv.")
    sheet = book.sheet_by_index(0)
    rows: list[list[Any]] = []
    for r in range(min(sheet.nrows, MAX_ROWS + 20)):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            else:
                row.append(cell.value)
        rows.append(row)
    return rows, sheet.name


def _kind(file_name: Optional[str], data: bytes) -> str:
    if data[:2] == b"PK":
        return "xlsx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    name = (file_name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if name.endswith(".xls"):
        return "xls"
    return "csv"


def read_sheet(data: bytes, file_name: Optional[str] = None) -> Sheet:
    if not data:
        raise ImportFileError("El archivo está vacío")
    if len(data) > MAX_BYTES:
        raise ImportFileError("El archivo supera 5 MB. Divídelo en partes y súbelas una por una.")
    kind = _kind(file_name, data)
    sheet_name = None
    if kind == "xlsx":
        raw_rows, sheet_name = _read_xlsx(data)
    elif kind == "xls":
        raw_rows, sheet_name = _read_xls(data)
    else:
        raw_rows = _read_csv(data)

    # Muchos programas ponen un título ("Listado de clientes") antes de los
    # encabezados: se toma como encabezado la primera fila con al menos dos
    # columnas reconocibles.
    header_idx = 0
    for i, row in enumerate(raw_rows[:15]):
        if sum(1 for c in row if detect_field(c)) >= 2:
            header_idx = i
            break
    if not raw_rows:
        return Sheet(headers=[], rows=[], sheet_name=sheet_name)
    headers = ["" if c is None else str(c).strip() for c in raw_rows[header_idx]]
    rows: list[tuple[int, list[Any]]] = []
    for offset, row in enumerate(raw_rows[header_idx + 1:], start=header_idx + 2):
        cells = list(row) + [None] * (len(headers) - len(row))
        if all(_blank(c) for c in cells):
            continue
        rows.append((offset, cells))
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"El archivo tiene más de {MAX_ROWS:,} filas. Divídelo en partes.".replace(",", "."))
    return Sheet(headers=headers, rows=rows, sheet_name=sheet_name)


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def auto_mapping(headers: list[str]) -> dict[int, Optional[str]]:
    mapping: dict[int, Optional[str]] = {}
    used: set[str] = set()
    for i, h in enumerate(headers):
        key = detect_field(h)
        if key and (key not in used or FIELD_BY_KEY[key].joinable):
            mapping[i] = key
            used.add(key)
        else:
            mapping[i] = None
    return mapping


# ─── Interpretación de valores ───────────────────────────────────────────────

class ValueProblem(ValueError):
    pass


def _text(v: Any) -> Optional[str]:
    if _blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


_PARTICLES = {"de", "del", "la", "las", "los", "y", "e", "da", "di", "van", "von"}


def smart_title(s: str) -> str:
    """"MARIA DE LOS ANGELES" → "Maria de los Angeles". Solo si viene todo en
    mayúsculas o minúsculas: un nombre escrito a mano se respeta."""
    letters = [c for c in s if c.isalpha()]
    if not letters or not (all(c.isupper() for c in letters) or all(c.islower() for c in letters)):
        return s
    words = s.lower().split(" ")
    return " ".join(w if (i and w in _PARTICLES) else w[:1].upper() + w[1:] for i, w in enumerate(words))


def normalize_document(v: Any) -> Optional[str]:
    s = _text(v)
    if not s:
        return None
    s = re.sub(r"[^0-9A-Za-z]", "", s).upper()
    return s or None


_ID_TYPES = {
    "cc": IdType.CC, "cedula": IdType.CC, "cedula de ciudadania": IdType.CC, "c c": IdType.CC,
    "ce": IdType.CE, "cedula de extranjeria": IdType.CE, "c e": IdType.CE,
    "ti": IdType.TI, "tarjeta de identidad": IdType.TI, "t i": IdType.TI,
    "pp": IdType.PP, "pasaporte": IdType.PP, "pa": IdType.PP,
    "nit": IdType.NIT,
}


def parse_id_type(v: Any) -> tuple[Optional[IdType], Optional[str]]:
    s = _norm(v)
    if not s:
        return None, None
    if s in _ID_TYPES:
        return _ID_TYPES[s], None
    if s in {"ppt", "pep", "permiso por proteccion temporal", "permiso especial de permanencia"}:
        return IdType.PP, "PPT/PEP se registró como pasaporte"
    raise ValueProblem(f"tipo de documento «{v}» no reconocido")


_GENDERS = {
    "m": Gender.M, "masculino": Gender.M, "hombre": Gender.M, "h": Gender.M, "male": Gender.M,
    "varon": Gender.M, "masc": Gender.M,
    "f": Gender.F, "femenino": Gender.F, "mujer": Gender.F, "female": Gender.F, "fem": Gender.F,
    "otro": Gender.other, "other": Gender.other, "no binario": Gender.other, "x": Gender.other,
}


def parse_gender(v: Any) -> Optional[Gender]:
    s = _norm(v)
    if not s:
        return None
    if s in _GENDERS:
        return _GENDERS[s]
    raise ValueProblem(f"sexo «{v}» no reconocido")


_DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y", "%Y%m%d")


def parse_date(v: Any) -> tuple[Optional[date], Optional[str]]:
    """Fechas de Colombia (día primero). Acepta celdas de fecha, el número
    serial de Excel y texto en los formatos usuales."""
    if _blank(v):
        return None, None
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 1000 < v < 80000:
            return date(1899, 12, 30) + timedelta(days=int(v)), None
        if 19000101 <= v <= 21001231:
            v = str(int(v))
        else:
            raise ValueProblem(f"fecha «{v}» no reconocida")
    s = str(v).strip().split("T")[0].split(" ")[0]
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000 if y + 2000 <= today_bogota().year else 1900
        warning = None
        if a <= 12 < b:
            a, b = b, a
            warning = f"«{s}» parecía mes/día: se leyó como {a:02d}/{b:02d}/{y}"
        try:
            return date(y, b, a), warning
        except ValueError:
            raise ValueProblem(f"fecha «{s}» no existe")
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    raise ValueProblem(f"fecha «{v}» no reconocida (usa dd/mm/aaaa)")


def parse_blood(v: Any) -> Optional[str]:
    if _blank(v):
        return None
    s = _norm(v).upper().replace(" ", "")
    raw = str(v).upper()
    sign = None
    if "+" in raw or "POS" in s:
        sign = "+"
    elif "-" in raw or "NEG" in s:
        sign = "-"
    letters = re.sub(r"POSITIVO|NEGATIVO|POS|NEG|RH|GRUPO", "", s)
    letters = letters.replace("0", "O")
    group = letters if letters in {"A", "B", "AB", "O"} else None
    if not group or not sign:
        raise ValueProblem(f"RH «{v}» no reconocido (ej. O+, A-)")
    return group + sign


def parse_phone(v: Any) -> Optional[str]:
    s = _text(v)
    if not s:
        return None
    digits = re.sub(r"\D", "", s)
    if not digits:
        return None
    if digits.startswith("57") and len(digits) == 12:
        digits = digits[2:]
    if len(digits) == 10 and digits[0] in "36":
        return f"+57{digits}"
    if len(digits) < 7:
        raise ValueProblem(f"teléfono «{s}» incompleto")
    return digits


_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


def parse_email(v: Any) -> Optional[str]:
    s = _text(v)
    if not s:
        return None
    s = s.lower().replace(" ", "")
    if not _EMAIL.match(s):
        raise ValueProblem(f"correo «{s}» inválido")
    return s


_TRUE = {"si", "s", "x", "1", "true", "verdadero", "yes", "y", "habilitado", "activo", "ok"}
_FALSE = {"no", "n", "0", "false", "falso", "inactivo", "deshabilitado"}


def parse_bool(v: Any) -> Optional[bool]:
    s = _norm(v)
    if not s:
        return None
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    raise ValueProblem(f"«{v}» no es sí/no")


_GOAL_WORDS: tuple[tuple[TrainingGoal, tuple[str, ...]], ...] = (
    (TrainingGoal.weight_loss, ("bajar", "perder", "adelgaz", "grasa", "peso ideal", "tonific", "definir", "reducir")),
    (TrainingGoal.muscle_gain, ("masa", "muscul", "hipertrof", "volumen", "aumentar", "ganar peso", "fuerza", "crecer")),
    (TrainingGoal.general_fitness, ("salud", "bienestar", "estar bien", "mantener", "acondicion", "general", "estres")),
    (TrainingGoal.endurance, ("resistencia", "cardio", "correr", "maraton", "aerobic", "condicion")),
    (TrainingGoal.flexibility, ("flexib", "estira", "movilidad", "yoga", "postura")),
)


def infer_goal(text: Optional[str]) -> Optional[TrainingGoal]:
    if not text:
        return None
    s = _norm(text)
    if s.replace(" ", "_") in TrainingGoal.__members__:
        return TrainingGoal(s.replace(" ", "_"))
    for goal, words in _GOAL_WORDS:
        if any(w in s for w in words):
            return goal
    return None


def parse_int(v: Any) -> Optional[int]:
    if _blank(v):
        return None
    try:
        n = int(float(str(v).replace(",", ".").strip()))
    except ValueError:
        raise ValueProblem(f"«{v}» no es un número")
    if n < 0:
        raise ValueProblem("no puede ser negativo")
    return n


# ─── Análisis por fila ───────────────────────────────────────────────────────

@dataclass
class ParsedRow:
    rows: list[int]
    fields: dict[str, Any] = field(default_factory=dict)
    trainer_name: Optional[str] = None
    plan_name: Optional[str] = None
    start: Optional[date] = None
    end: Optional[date] = None
    entries: Optional[int] = None
    warnings: list[str] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)

    @property
    def doc(self) -> Optional[str]:
        return self.fields.get("id_number")

    def key(self) -> Optional[tuple]:
        if self.doc:
            return ("doc", self.doc)
        name = _norm(self.fields.get("full_name"))
        phone = (self.fields.get("phone") or "")[-10:]
        if name and phone:
            return ("name_phone", name, phone)
        if name and self.fields.get("email"):
            return ("name_email", name, self.fields["email"])
        return None


def _joined(cells: list[Any], cols: list[int], sep: str = " ") -> Any:
    vals = [cells[c] for c in cols if c < len(cells) and not _blank(cells[c])]
    if not vals:
        return None
    if len(vals) == 1:
        return vals[0]
    return sep.join(_text(v) or "" for v in vals)


def parse_row(line: int, cells: list[Any], mapping: dict[int, Optional[str]]) -> ParsedRow:
    by_field: dict[str, list[int]] = {}
    for col, key in sorted(mapping.items()):
        if key:
            by_field.setdefault(key, []).append(col)
    out = ParsedRow(rows=[line])

    def get(key: str) -> Any:
        cols = by_field.get(key)
        if not cols:
            return None
        return _joined(cells, cols, sep=" " if key != "notes" else " · ")

    def attempt(key: str, fn):
        try:
            return fn(get(key))
        except ValueProblem as exc:
            out.warnings.append(f"{FIELD_BY_KEY[key].label}: {exc} (se dejó vacío)")
            return None

    name = _text(get("full_name"))
    if name:
        out.fields["full_name"] = smart_title(name)[:255]

    raw_doc = get("id_number")
    doc = normalize_document(raw_doc)
    if doc and len(doc) < 3:
        out.warnings.append(f"Cédula «{_text(raw_doc)}» demasiado corta (se dejó sin documento)")
        doc = None
    if doc:
        out.fields["id_number"] = doc[:50]
        if "id_type" in by_field:
            try:
                id_type, note = parse_id_type(get("id_type"))
                if note:
                    out.warnings.append(note)
                if id_type:
                    out.fields["id_type"] = id_type
            except ValueProblem as exc:
                out.warnings.append(f"Tipo de documento: {exc}; se asumió cédula de ciudadanía")

    for key, fn in (("gender", parse_gender), ("email", parse_email), ("phone", parse_phone),
                    ("emergency_contact_phone", parse_phone), ("blood_type", parse_blood),
                    ("access_by_document", parse_bool)):
        if key in by_field:
            val = attempt(key, fn)
            if val is not None:
                out.fields[key] = val

    if "birth_date" in by_field:
        try:
            bd, note = parse_date(get("birth_date"))
            if note:
                out.warnings.append(f"Fecha de nacimiento: {note}")
            if bd and not (date(1900, 1, 1) <= bd <= today_bogota()):
                out.warnings.append(f"Fecha de nacimiento {bd:%d/%m/%Y} fuera de rango (se dejó vacía)")
                bd = None
            if bd:
                out.fields["birth_date"] = bd
        except ValueProblem as exc:
            out.warnings.append(f"Fecha de nacimiento: {exc} (se dejó vacía)")

    for key, limit in (("city", 100), ("eps", 100), ("address", 500), ("emergency_contact_name", 255),
                       ("notes", 1000)):
        val = _text(get(key))
        if val:
            out.fields[key] = (smart_title(val) if key in ("city", "emergency_contact_name") else val)[:limit]

    country = _text(get("country"))
    if country:
        c = _norm(country)
        out.fields["country"] = "CO" if c in {"co", "col", "colombia"} else country.upper()[:2]

    expectation = _text(get("expectation"))
    if expectation:
        goal = infer_goal(expectation)
        if _norm(expectation).replace(" ", "_") not in TrainingGoal.__members__:
            out.fields["expectation"] = expectation[:500]
        if goal:
            out.fields["training_goal"] = goal

    out.trainer_name = _text(get("trainer"))
    out.plan_name = _text(get("plan"))
    for key, attr in (("membership_start", "start"), ("membership_end", "end")):
        if key in by_field:
            try:
                d, note = parse_date(get(key))
                if note:
                    out.warnings.append(f"{FIELD_BY_KEY[key].label}: {note}")
                setattr(out, attr, d)
            except ValueProblem as exc:
                out.warnings.append(f"{FIELD_BY_KEY[key].label}: {exc}")
    if "entries_remaining" in by_field:
        out.entries = attempt("entries_remaining", parse_int)

    if not out.fields.get("full_name"):
        out.errors.append({"field": "full_name", "message": "Falta el nombre"})
    elif len(out.fields["full_name"]) < 2:
        out.errors.append({"field": "full_name", "message": "Nombre demasiado corto"})
    return out


def _merge(into: ParsedRow, other: ParsedRow) -> None:
    """Dos filas de la misma persona: la primera manda, la segunda completa."""
    into.rows.extend(other.rows)
    for k, v in other.fields.items():
        into.fields.setdefault(k, v)
    into.trainer_name = into.trainer_name or other.trainer_name
    into.plan_name = into.plan_name or other.plan_name
    if other.end and (not into.end or other.end > into.end):
        into.start, into.end, into.entries = other.start, other.end, other.entries
        into.plan_name = other.plan_name or into.plan_name
    into.warnings.extend(other.warnings)
    into.errors.extend(other.errors)


# ─── Clasificación contra la base ────────────────────────────────────────────

@dataclass
class Options:
    mapping: dict[int, Optional[str]]
    on_conflict: str = "keep"  # keep | file
    plan_map: dict[str, Optional[str]] = field(default_factory=dict)
    trainer_map: dict[str, Optional[str]] = field(default_factory=dict)


@dataclass
class Plan:
    """Lo que se va a hacer con una persona del archivo."""
    parsed: ParsedRow
    action: str  # create | update | unchanged | error
    member: Optional[Member] = None
    values: dict[str, Any] = field(default_factory=dict)       # para crear
    fill: dict[str, Any] = field(default_factory=dict)         # campos vacíos a llenar
    conflicts: list[dict] = field(default_factory=list)        # existente ≠ archivo
    membership: Optional[dict] = None
    make_member_type: bool = False


def _display(v: Any) -> Any:
    if isinstance(v, (date, datetime)):
        return v.strftime("%d/%m/%Y")
    if hasattr(v, "value"):
        return v.value
    if isinstance(v, uuid.UUID):
        return str(v)
    return v


def _same(a: Any, b: Any, key: str = "") -> bool:
    if key == "id_number":
        return normalize_document(a) == normalize_document(b)
    if key in ("phone", "emergency_contact_phone"):
        return re.sub(r"\D", "", str(a))[-10:] == re.sub(r"\D", "", str(b))[-10:]
    if isinstance(a, str) and isinstance(b, str):
        return _norm(a) == _norm(b)
    return getattr(a, "value", a) == getattr(b, "value", b)


class _Catalog:
    def __init__(self, members: list[Member], plans: list[MemberPlan], trainers: list[Instructor],
                 memberships: dict[uuid.UUID, list[Membership]]):
        self.members = members
        self.plans = plans
        self.trainers = trainers
        self.memberships = memberships
        self.by_doc: dict[str, Member] = {}
        self.by_email: dict[str, Member] = {}
        self.by_name_phone: dict[tuple, Member] = {}
        self.by_name_email: dict[tuple, Member] = {}
        for m in members:
            if m.id_number and (d := normalize_document(m.id_number)):
                self.by_doc.setdefault(d, m)
            if m.email:
                self.by_email.setdefault(m.email.lower(), m)
            name = _norm(m.full_name)
            if m.phone:
                self.by_name_phone.setdefault((name, re.sub(r"\D", "", m.phone)[-10:]), m)
            if m.email:
                self.by_name_email.setdefault((name, m.email.lower()), m)
        self.plan_by_name = {_norm(p.name): p for p in plans}
        self.plan_by_id = {str(p.id): p for p in plans}
        self.trainer_by_id = {str(t.id): t for t in trainers}

    def find_member(self, p: ParsedRow) -> Optional[Member]:
        if p.doc:
            return self.by_doc.get(p.doc)
        name = _norm(p.fields.get("full_name"))
        phone = (p.fields.get("phone") or "")
        if phone:
            hit = self.by_name_phone.get((name, re.sub(r"\D", "", phone)[-10:]))
            if hit:
                return hit
        if p.fields.get("email"):
            return self.by_name_email.get((name, p.fields["email"]))
        return None

    def match_plan(self, raw: Optional[str], opts: Options) -> tuple[Optional[MemberPlan], bool]:
        """(plan, ¿se resolvió?). Un nombre sin plan equivalente queda sin resolver."""
        if not raw:
            return None, True
        key = _norm(raw)
        if key in opts.plan_map:
            chosen = opts.plan_map[key]
            return (self.plan_by_id.get(chosen) if chosen else None), True
        if key in self.plan_by_name:
            return self.plan_by_name[key], True
        return None, False

    def match_trainer(self, raw: Optional[str], opts: Options) -> tuple[Optional[Instructor], bool]:
        if not raw:
            return None, True
        key = _norm(raw)
        if key in opts.trainer_map:
            chosen = opts.trainer_map[key]
            return (self.trainer_by_id.get(chosen) if chosen else None), True
        tokens = set(key.split())
        exact = [t for t in self.trainers if _norm(t.full_name) == key]
        if len(exact) == 1:
            return exact[0], True
        partial = [t for t in self.trainers
                   if tokens and (tokens <= set(_norm(t.full_name).split()) or set(_norm(t.full_name).split()) <= tokens)]
        if len(partial) == 1:
            return partial[0], True
        return None, False


async def _load_catalog(db: AsyncSession) -> _Catalog:
    members = list((await db.execute(select(Member))).scalars().all())
    plans = list((await db.execute(select(MemberPlan).where(MemberPlan.is_active.is_(True)))).scalars().all())
    trainers = list((await db.execute(select(Instructor).where(Instructor.is_active.is_(True)))).scalars().all())
    by_member: dict[uuid.UUID, list[Membership]] = {}
    for ms in (await db.execute(select(Membership))).scalars().all():
        by_member.setdefault(ms.member_id, []).append(ms)
    return _Catalog(members, plans, trainers, by_member)


def _membership_plan(p: ParsedRow, cat: _Catalog, opts: Options, unmatched_plans: dict[str, int]) -> Optional[dict]:
    if not (p.end or p.start or p.plan_name or p.entries is not None):
        return None
    plan, resolved = cat.match_plan(p.plan_name, opts)
    if not resolved:
        unmatched_plans[p.plan_name] = unmatched_plans.get(p.plan_name, 0) + 1
    duration = plan.duration_days if plan else None
    start, end = p.start, p.end
    if end and not start:
        start = end - timedelta(days=duration or 30)
    elif start and not end:
        if not duration:
            p.warnings.append("Plan sin fecha de vencimiento ni duración conocida: no se creó membresía")
            return None
        end = start + timedelta(days=duration)
    elif not start and not end:
        if p.entries is None:
            p.warnings.append("Plan sin fechas: no se creó membresía (agrega la fecha de vencimiento)")
            return None
        start = today_bogota()
        end = start + timedelta(days=duration or 30)
    if end < start:
        p.warnings.append("La fecha de vencimiento es anterior al inicio: no se creó membresía")
        return None
    entries = p.entries
    if entries is None and plan is not None and plan.entries_total:
        entries = plan.entries_total
    return {
        "plan": plan,
        "plan_name": (plan.name if plan else (smart_title(p.plan_name) if p.plan_name else IMPORTED_PLAN_NAME))[:100],
        "start": start,
        "end": end,
        "entries": entries,
        "unresolved_plan": not resolved,
    }


def _validate(values: dict, p: ParsedRow) -> list[dict]:
    """Valida con el mismo schema del alta. Un dato opcional inválido no
    tumba la fila: se descarta con aviso. Solo el nombre es indispensable."""
    for _ in range(3):
        try:
            MemberBase(**{k: v for k, v in values.items() if k in MemberBase.model_fields})
            return []
        except ValidationError as exc:
            bad = {str(e["loc"][0]) for e in exc.errors() if e["loc"]}
            if not bad or "full_name" in bad:
                return [{"field": str(e["loc"][0]) if e["loc"] else "fila", "message": e["msg"]} for e in exc.errors()]
            for k in bad:
                label = FIELD_BY_KEY[k].label if k in FIELD_BY_KEY else k
                p.warnings.append(f"{label}: valor inválido, se dejó vacío")
                values.pop(k, None)
                p.fields.pop(k, None)
    return [{"field": "fila", "message": "Datos inválidos"}]


def build_plans(sheet: Sheet, opts: Options, cat: _Catalog) -> tuple[list[Plan], dict, dict]:
    # 1) Interpretar cada fila y unir las que son la misma persona dentro del archivo.
    people: list[ParsedRow] = []
    by_key: dict[tuple, ParsedRow] = {}
    for line, cells in sheet.rows:
        parsed = parse_row(line, cells, opts.mapping)
        k = parsed.key() if not parsed.errors else None
        if k and k in by_key:
            first = by_key[k]
            _merge(first, parsed)
            first.warnings.append(f"La fila {line} es la misma persona: se unieron")
            continue
        if k:
            by_key[k] = parsed
        people.append(parsed)

    unmatched_plans: dict[str, int] = {}
    unmatched_trainers: dict[str, int] = {}
    emails_taken: dict[str, int] = {}  # correo → id(persona del archivo que lo tomó)
    plans: list[Plan] = []

    for p in people:
        if p.errors:
            plans.append(Plan(parsed=p, action="error"))
            continue

        existing = cat.find_member(p)

        # Correo compartido: si lo tiene otra ficha o lo tomó otra persona del archivo
        # antes, esta persona entra sin correo.
        email = p.fields.get("email")
        if email:
            owner = cat.by_email.get(email)
            if (owner is not None and owner is not existing) or emails_taken.get(email, id(p)) != id(p):
                p.fields.pop("email")
                p.warnings.append(f"El correo {email} ya lo usa otra persona: se importó sin correo")
            else:
                emails_taken[email] = id(p)

        trainer, resolved = cat.match_trainer(p.trainer_name, opts)
        if not resolved:
            unmatched_trainers[p.trainer_name] = unmatched_trainers.get(p.trainer_name, 0) + 1
            p.warnings.append(f"Entrenador «{p.trainer_name}» no existe en LSMFit: se dejó sin entrenador")
        if trainer:
            p.fields["trainer_id"] = trainer.id

        if not p.doc:
            p.warnings.append("Sin documento: complétalo después para que pueda entrar con cédula o al torniquete")

        ms = _membership_plan(p, cat, opts, unmatched_plans)

        if existing is None:
            values = dict(p.fields)
            values.setdefault("country", "CO")
            if values.get("id_number"):
                values.setdefault("id_type", IdType.CC)
            problems = _validate(values, p)
            if problems:
                p.errors.extend(problems)
                plans.append(Plan(parsed=p, action="error"))
                continue
            plans.append(Plan(parsed=p, action="create", values=values, membership=ms))
            continue

        plan = Plan(parsed=p, action="update", member=existing)
        for k, incoming in p.fields.items():
            current = getattr(existing, k, None)
            if k == "access_by_document":
                if incoming and not current:
                    plan.fill[k] = True
                continue
            if current in (None, ""):
                plan.fill[k] = incoming
            elif not _same(current, incoming, k):
                plan.conflicts.append({
                    "field": k,
                    "label": FIELD_BY_KEY[k].label if k in FIELD_BY_KEY else k,
                    "current": _display(current),
                    "incoming": _display(incoming),
                })
        if existing.member_type == MemberType.visitor.value:
            plan.make_member_type = True
        if ms:
            current_ms = cat.memberships.get(existing.id, [])
            if any(m.end_date == ms["end"] and m.status != MembershipStatus.cancelled.value for m in current_ms):
                ms = None  # ya importada (o ya registrada igual): no se duplica
            elif any(m.end_date >= ms["end"] and m.status != MembershipStatus.cancelled.value for m in current_ms):
                p.warnings.append("Ya tiene en LSMFit una membresía que vence igual o después: no se agregó la del archivo")
                ms = None
        plan.membership = ms
        if not plan.fill and not ms and not plan.make_member_type and not (plan.conflicts and opts.on_conflict == "file"):
            plan.action = "unchanged"
        plans.append(plan)

    return plans, unmatched_plans, unmatched_trainers


def _summary(plans: list[Plan]) -> dict:
    s = {"create": 0, "update": 0, "unchanged": 0, "error": 0, "memberships": 0, "conflicts": 0,
         "warnings": 0, "rows": sum(len(p.parsed.rows) for p in plans)}
    for p in plans:
        s[p.action] += 1
        s["memberships"] += 1 if p.membership and p.action != "error" else 0
        s["conflicts"] += len(p.conflicts)
        s["warnings"] += len(p.parsed.warnings)
    return s


def _row_out(p: Plan) -> dict:
    ms = p.membership
    return {
        "rows": p.parsed.rows,
        "name": p.parsed.fields.get("full_name") or (p.member.full_name if p.member else None),
        "id_number": p.parsed.doc,
        "action": p.action,
        "fill": sorted(FIELD_BY_KEY[k].label if k in FIELD_BY_KEY else k for k in p.fill if k != "trainer_id")
                + (["Entrenador"] if "trainer_id" in p.fill else []),
        "conflicts": p.conflicts,
        "membership": {
            "plan_name": ms["plan_name"], "start": ms["start"].isoformat(), "end": ms["end"].isoformat(),
            "entries": ms["entries"],
        } if ms else None,
        "warnings": p.parsed.warnings,
        "errors": p.parsed.errors,
    }


def _parse_options(sheet: Sheet, raw: Optional[dict]) -> Options:
    raw = raw or {}
    mapping = auto_mapping(sheet.headers)
    for k, v in (raw.get("mapping") or {}).items():
        try:
            idx = int(k)
        except (TypeError, ValueError):
            continue
        if 0 <= idx < len(sheet.headers):
            mapping[idx] = v if v in FIELD_BY_KEY else None
    on_conflict = raw.get("on_conflict") if raw.get("on_conflict") in ("keep", "file") else "keep"
    return Options(
        mapping=mapping,
        on_conflict=on_conflict,
        plan_map={_norm(k): v for k, v in (raw.get("plan_map") or {}).items()},
        trainer_map={_norm(k): v for k, v in (raw.get("trainer_map") or {}).items()},
    )


async def preview(db: AsyncSession, data: bytes, file_name: Optional[str], raw_options: Optional[dict]) -> dict:
    # Leer un Excel grande es CPU puro: fuera del event loop para no congelar la API.
    sheet = await asyncio.to_thread(read_sheet, data, file_name)
    opts = _parse_options(sheet, raw_options)
    cat = await _load_catalog(db)
    plans, unmatched_plans, unmatched_trainers = build_plans(sheet, opts, cat)

    columns = []
    for i, h in enumerate(sheet.headers):
        samples = [_display(_text(cells[i])) for _, cells in sheet.rows[:50] if i < len(cells) and not _blank(cells[i])][:3]
        columns.append({"index": i, "header": h, "field": opts.mapping.get(i), "samples": samples})

    # Primero lo que necesita atención (errores, conflictos, avisos), luego lo demás.
    ordered = sorted(plans, key=lambda p: (
        p.action != "error", not p.conflicts, not p.parsed.warnings, p.parsed.rows[0]))
    return {
        "file_name": file_name,
        "sheet": sheet.sheet_name,
        "columns": columns,
        "fields": [{"key": f.key, "label": f.label} for f in FIELDS],
        "summary": _summary(plans),
        "rows": [_row_out(p) for p in ordered[:PREVIEW_ROWS]],
        "truncated": len(plans) > PREVIEW_ROWS,
        "unmatched_plans": [{"name": k, "count": v} for k, v in sorted(unmatched_plans.items())],
        "unmatched_trainers": [{"name": k, "count": v} for k, v in sorted(unmatched_trainers.items())],
        "plans": [{"id": str(p.id), "name": p.name, "duration_days": p.duration_days} for p in cat.plans],
        "trainers": [{"id": str(t.id), "name": t.full_name} for t in cat.trainers],
        "on_conflict": opts.on_conflict,
    }


def _new_membership(member_id: uuid.UUID, ms: dict, batch_id: uuid.UUID) -> Membership:
    plan: Optional[MemberPlan] = ms["plan"]
    return Membership(
        member_id=member_id,
        plan_id=plan.id if plan else None,
        plan_name=ms["plan_name"],
        access_type=getattr(plan, "access_type", None) or "gym_and_classes",
        class_credits_per_month=getattr(plan, "class_credits_per_month", None),
        includes_trainer=bool(getattr(plan, "includes_trainer", False)),
        trainer_sessions_per_month=getattr(plan, "trainer_sessions_per_month", None),
        # Ya se cobró en el sistema anterior: precio 0 y sin pago, para no
        # inflar ingresos ni aparecer como deuda.
        price=0,
        start_date=ms["start"],
        end_date=ms["end"],
        duration_days=max((ms["end"] - ms["start"]).days, 1),
        status=MembershipStatus.active.value,
        notes="Importada del sistema anterior",
        source="import",
        import_batch_id=batch_id,
        entries_total=ms["entries"],
        entries_used=0,
    )


def _json_value(v: Any) -> Any:
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, uuid.UUID):
        return str(v)
    return getattr(v, "value", v)


async def apply(db: AsyncSession, data: bytes, file_name: Optional[str], raw_options: Optional[dict],
                actor_id: Optional[uuid.UUID]) -> dict:
    sheet = await asyncio.to_thread(read_sheet, data, file_name)
    opts = _parse_options(sheet, raw_options)
    cat = await _load_catalog(db)
    plans, _, _ = build_plans(sheet, opts, cat)

    batch = ImportBatch(kind="members", file_name=(file_name or "")[:255] or None, status="applied",
                        created_by=actor_id)
    db.add(batch)
    await db.flush()

    changes: list[dict] = []
    for plan in plans:
        if plan.action == "create":
            # Sin consent_data_at: la autorización del sistema anterior no consta aquí.
            # Se le pide al socio con el enlace de completar datos (member_profile).
            member = Member(**plan.values, import_batch_id=batch.id)
            db.add(member)
            await db.flush()
            entry: dict[str, Any] = {"member_id": str(member.id), "created": True, "memberships": []}
            if plan.membership:
                ms = _new_membership(member.id, plan.membership, batch.id)
                db.add(ms)
                await db.flush()
                entry["memberships"].append(str(ms.id))
            changes.append(entry)
        elif plan.action == "update":
            member = plan.member
            touched: dict[str, list] = {}
            for k, v in plan.fill.items():
                touched[k] = [_json_value(getattr(member, k)), _json_value(v)]
                setattr(member, k, v)
            if opts.on_conflict == "file":
                for c in plan.conflicts:
                    k = c["field"]
                    v = plan.parsed.fields[k]
                    touched[k] = [_json_value(getattr(member, k)), _json_value(v)]
                    setattr(member, k, v)
            if plan.make_member_type:
                touched["member_type"] = [member.member_type, MemberType.member.value]
                member.member_type = MemberType.member.value
            entry = {"member_id": str(member.id), "created": False, "fields": touched, "memberships": []}
            if plan.membership:
                ms = _new_membership(member.id, plan.membership, batch.id)
                db.add(ms)
                await db.flush()
                entry["memberships"].append(str(ms.id))
            changes.append(entry)

    summary = _summary(plans)
    errors = [
        {"row": p.parsed.rows[0], "field": e["field"], "message": e["message"]}
        for p in plans if p.action == "error" for e in p.parsed.errors
    ]
    batch.counts = {**summary, "errors": errors[:500]}
    batch.changes = changes
    await db.flush()
    return {"batch_id": str(batch.id), "summary": summary, "errors": errors}


# ─── Deshacer un lote ───────────────────────────────────────────────────────

async def _member_has_money(db: AsyncSession, member_id: uuid.UUID) -> bool:
    pagos = await db.scalar(select(func.count()).select_from(Payment).where(Payment.member_id == member_id))
    pases = await db.scalar(select(func.count()).select_from(DayPass).where(DayPass.member_id == member_id))
    return bool(pagos or pases)


async def _delete_memberships(db: AsyncSession, ids: Iterable[uuid.UUID]) -> tuple[int, int]:
    """Borra membresías importadas que no tengan cobros asociados. (borradas, conservadas)."""
    removed = kept = 0
    for mid in ids:
        ms = await db.get(Membership, mid)
        if ms is None:
            continue
        paid = await db.scalar(select(func.count()).select_from(Payment).where(Payment.membership_id == mid))
        if paid:
            kept += 1
            continue
        await db.execute(delete(MembershipFreeze).where(MembershipFreeze.membership_id == mid))
        await db.delete(ms)
        removed += 1
    return removed, kept


async def undo(db: AsyncSession, batch: ImportBatch, actor_id: Optional[uuid.UUID]) -> dict:
    if batch.status == "undone":
        raise ValueError("Este lote ya se deshizo")
    result = {"members_removed": 0, "members_kept": 0, "fields_restored": 0, "fields_kept": 0,
              "memberships_removed": 0, "memberships_kept": 0}
    for entry in reversed(batch.changes or []):
        member_id = uuid.UUID(entry["member_id"])
        removed, kept = await _delete_memberships(db, [uuid.UUID(x) for x in entry.get("memberships", [])])
        result["memberships_removed"] += removed
        result["memberships_kept"] += kept
        member = await db.get(Member, member_id)
        if member is None:
            continue
        if entry.get("created"):
            # Lo que el gym registró después (cobros) no se borra: la ficha se conserva.
            if await _member_has_money(db, member_id):
                result["members_kept"] += 1
                continue
            await db.execute(delete(MembershipFreeze).where(MembershipFreeze.membership_id.in_(
                select(Membership.id).where(Membership.member_id == member_id))))
            for model in (Membership, CheckIn, ClassEnrollment, SessionBooking, MemberRoutine):
                await db.execute(delete(model).where(model.member_id == member_id))
            await db.delete(member)
            result["members_removed"] += 1
            continue
        for k, (old, new) in (entry.get("fields") or {}).items():
            # Solo se restaura si nadie lo cambió después de importar.
            if _json_value(getattr(member, k, None)) == new:
                restored = old
                if old is not None and k == "birth_date":
                    restored = date.fromisoformat(old)
                elif old is not None and k == "trainer_id":
                    restored = uuid.UUID(old)
                setattr(member, k, restored)
                result["fields_restored"] += 1
            else:
                result["fields_kept"] += 1
    batch.status = "undone"
    batch.undone_at = utcnow_naive()
    batch.undone_by = actor_id
    batch.counts = {**(batch.counts or {}), "undo": result}
    await db.flush()
    return result


# ─── Plantilla ──────────────────────────────────────────────────────────────

TEMPLATE_COLUMNS: tuple[tuple[str, str, bool], ...] = (
    ("id_number", "Cédula", False),
    ("id_type", "Tipo de documento", False),
    ("full_name", "Nombre", True),
    ("gender", "Sexo", False),
    ("birth_date", "Fecha Nac", False),
    ("city", "Ciudad", False),
    ("email", "Email", False),
    ("phone", "Teléfono", False),
    ("eps", "EPS", False),
    ("blood_type", "RH", False),
    ("expectation", "Expectativa", False),
    ("address", "Dirección", False),
    ("access_by_document", "Ingreso por DNI", False),
    ("trainer", "Entrenador", False),
    ("emergency_contact_name", "Contacto de emergencia", False),
    ("emergency_contact_phone", "Teléfono de emergencia", False),
    ("plan", "Plan", False),
    ("membership_start", "Fecha inicio", False),
    ("membership_end", "Fecha vencimiento", False),
    ("entries_remaining", "Entradas restantes", False),
    ("notes", "Notas", False),
)

_EXAMPLE = {
    "id_number": "1018456789", "id_type": "CC", "full_name": "María Gómez", "gender": "Femenino",
    "birth_date": "12/04/1995", "city": "Bogotá", "email": "maria@ejemplo.co", "phone": "300 123 4567",
    "eps": "Sura", "blood_type": "O+", "expectation": "Bajar de peso y tonificar",
    "address": "Calle 10 # 5-20", "access_by_document": "Sí", "trainer": "",
    "emergency_contact_name": "Juan Gómez", "emergency_contact_phone": "301 987 6543",
    "plan": "Mensualidad", "membership_start": "01/10/2026", "membership_end": "31/10/2026",
    "entries_remaining": "", "notes": "",
}


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Miembros"
    header_font = Font(bold=True, color="000000")
    header_fill = PatternFill("solid", fgColor=_BRAND_TINT)
    for col_idx, (key, label, required) in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label + (" *" if required else ""))
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = max(len(label) + 4, 16)
        ws.cell(row=2, column=col_idx, value=_EXAMPLE.get(key, ""))

    info = wb.create_sheet("Instrucciones")
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 80
    info.append(["Campo", "Cómo llenarlo"])
    for c in info[1]:
        c.font = Font(bold=True)
    for row in (
        ["Obligatorio", "Solo el Nombre. Todo lo demás se puede completar después."],
        ["Cédula", "Con o sin puntos. Es la llave: si la persona ya existe, solo se llenan los datos que falten."],
        ["Tipo de documento", "CC, CE, TI, PP o NIT. Vacío = CC."],
        ["Sexo", "M/F, Masculino/Femenino, Hombre/Mujer."],
        ["Fechas", "dd/mm/aaaa (ej. 12/04/1995)."],
        ["Teléfono", "Con o sin +57, espacios o guiones."],
        ["RH", "O+, A-, AB+... o 'O positivo'."],
        ["Ingreso por DNI", "Sí/No: puede entrar digitando la cédula."],
        ["Entrenador", "Nombre del entrenador como está en LSMFit."],
        ["Plan y fechas", "Para que el socio pueda entrar desde el primer día: nombre del plan y fecha de vencimiento."],
        ["Entradas restantes", "Solo para ticketeras: entradas que le quedan."],
        ["Otros encabezados", "También se reconocen los del programa anterior (Nombres, Apellidos, Celular, Vence...)."],
        ["Fila de ejemplo", "Bórrala antes de subir el archivo."],
    ):
        info.append(row)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
